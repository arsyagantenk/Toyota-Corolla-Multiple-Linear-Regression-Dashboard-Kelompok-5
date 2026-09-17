import io
import html
import re
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
import statsmodels.api as sm
from statsmodels.stats.outliers_influence import variance_inflation_factor

from data_utils import load_dataset
from preprocessing import preprocess, prepare_regression_split, transform_new_data
from modelling import train_ols, compute_metrics, extract_coefficients


DISPLAY_LABELS = {
    "Price": "Harga Mobil",
    "Age": "Usia Mobil",
    "Kilometers": "Jarak Tempuh",
    "Fuel Type": "Jenis Bahan Bakar",
    "Horse Power": "Tenaga Mesin",
    "Metallic": "Warna Metalik",
    "Automatic": "Transmisi Otomatis",
    "CC": "Kapasitas Mesin (CC)",
    "Doors": "Jumlah Pintu",
    "Quart Tax": "Pajak Tahunan",
    "Weight": "Bobot Kendaraan",
    "Model": "Model Corolla",
    "Id": "ID Kendaraan",
}

# Hanya variabel yang ditetapkan pada materi/tugas yang digunakan untuk analisis.
ANALYSIS_FEATURES = [
    "Age", "Kilometers", "Fuel Type", "Horse Power", "Metallic",
    "Automatic", "CC", "Doors", "Quart Tax", "Weight"
]

EXTREME_NUMERIC_FEATURES = [
    "Price", "Age", "Kilometers", "Horse Power", "CC", "Quart Tax", "Weight"
]


def extreme_value_analysis(data):
    """Detect extreme values using the 1st and 99th percentiles for each numeric variable."""
    summary_rows = []
    observation_rows = []
    any_extreme = pd.Series(False, index=data.index)

    for col in EXTREME_NUMERIC_FEATURES:
        if col not in data.columns:
            continue
        series = pd.to_numeric(data[col], errors="coerce")
        p1 = float(series.quantile(0.01))
        p99 = float(series.quantile(0.99))
        lower_mask = series < p1
        upper_mask = series > p99
        any_extreme = any_extreme | lower_mask.fillna(False) | upper_mask.fillna(False)

        summary_rows.append({
            "Variabel": display_name(col),
            "P1 (1%)": p1,
            "P99 (99%)": p99,
            "Lower Extreme (<P1)": int(lower_mask.sum()),
            "Upper Extreme (>P99)": int(upper_mask.sum()),
            "Total Extreme": int((lower_mask | upper_mask).sum()),
        })

        for idx in data.index[lower_mask.fillna(False)]:
            observation_rows.append({
                "ID": data.loc[idx, "Id"] if "Id" in data.columns else idx,
                "Variabel": display_name(col),
                "Nilai": float(series.loc[idx]),
                "Jenis Extreme": "Lower extreme (< P1)",
                "P1": p1,
                "P99": p99,
            })
        for idx in data.index[upper_mask.fillna(False)]:
            observation_rows.append({
                "ID": data.loc[idx, "Id"] if "Id" in data.columns else idx,
                "Variabel": display_name(col),
                "Nilai": float(series.loc[idx]),
                "Jenis Extreme": "Upper extreme (> P99)",
                "P1": p1,
                "P99": p99,
            })

    summary = pd.DataFrame(summary_rows)
    observations = pd.DataFrame(observation_rows)
    if not observations.empty:
        observations = observations.sort_values(["ID", "Variabel"]).reset_index(drop=True)
    return summary, observations, any_extreme


def influence_analysis(model, row_indices, raw_data):
    """Return leverage and Cook's Distance for observations used to fit the OLS model."""
    influence = model.get_influence()
    frame = influence.summary_frame().copy()
    frame["row_index"] = np.asarray(row_indices)
    if "Id" in raw_data.columns:
        frame["ID"] = [raw_data.loc[i, "Id"] if i in raw_data.index else i for i in frame["row_index"]]
    else:
        frame["ID"] = frame["row_index"]
    frame["Harga Aktual"] = [raw_data.loc[i, "Price"] if i in raw_data.index else np.nan for i in frame["row_index"]]
    frame["Leverage"] = frame["hat_diag"]
    frame["Cook's Distance"] = frame["cooks_d"]
    cols = ["ID", "Harga Aktual", "Leverage", "Cook's Distance", "standard_resid"]
    result = frame[cols].copy().sort_values("Cook's Distance", ascending=False).reset_index(drop=True)
    result.rename(columns={"standard_resid": "Standardized Residual"}, inplace=True)
    return result


def display_name(name):
    """Show dataset variable names rather than sklearn preprocessing names."""
    if name is None:
        return ""
    name = str(name)
    if name in {"const", "Intercept"}:
        return "Intercept"
    raw = name
    if "__" in raw:
        raw = raw.split("__", 1)[1]
    for base, label in DISPLAY_LABELS.items():
        if raw == base:
            return label
        prefix = base + "_"
        if raw.startswith(prefix):
            return f"{label}: {raw[len(prefix):].replace('_', ' ')}"
    return raw.replace("_", " ")



def build_vif_table(model):
    X = model.model.exog
    names = list(model.model.exog_names)
    rows = []
    for i, name in enumerate(names):
        if name == "const":
            continue
        try:
            vif = float(variance_inflation_factor(X, i))
        except Exception:
            vif = np.nan
        rows.append({"Feature": name, "VIF": vif})
    return pd.DataFrame(rows)


def build_model_summary_table(model, train_metrics, test_metrics, full_metrics=None):
    s = float(np.sqrt(model.ssr / model.df_resid)) if model.df_resid > 0 else np.nan
    rows = [
        {"Ukuran": "S (Standard Error of Regression) · Training", "Nilai": s},
        {"Ukuran": "R² Training", "Nilai": train_metrics["r2"]},
        {"Ukuran": "Adjusted R² Training", "Nilai": train_metrics["adj_r2"]},
        {"Ukuran": "R² Test", "Nilai": test_metrics["r2"]},
    ]
    if full_metrics is not None:
        rows.extend([
            {"Ukuran": "R² Full Model", "Nilai": full_metrics["r2"]},
            {"Ukuran": "Adjusted R² Full Model", "Nilai": full_metrics["adj_r2"]},
        ])
    rows.extend([
        {"Ukuran": "MAE Test", "Nilai": test_metrics["mae"]},
        {"Ukuran": "RMSE Test", "Nilai": test_metrics["rmse"]},
        {"Ukuran": "MAPE Test", "Nilai": test_metrics["mape"]},
    ])
    return pd.DataFrame(rows)


def fit_full_model_metrics(bundle):
    """Fit OLS on all modeling observations for direct comparison with full-data software output."""
    X_all = bundle["X_all"].copy()
    y_all = pd.Series(bundle["y_all"]).astype(float)
    full_model = sm.OLS(y_all, sm.add_constant(X_all, has_constant="add")).fit()
    pred_all = full_model.predict(sm.add_constant(X_all, has_constant="add"))
    metrics = compute_metrics(y_all, pred_all, len(bundle["feature_names"]))
    metrics["s"] = float(np.sqrt(full_model.ssr / full_model.df_resid)) if full_model.df_resid > 0 else np.nan
    return full_model, metrics


def build_anova_table(model, X_train, y_train, original_features):
    X_train = X_train.copy()
    y = pd.Series(y_train).astype(float)
    full_sse = float(model.ssr)
    mse = float(model.mse_resid)
    sst = float(np.sum((y - y.mean()) ** 2))
    rows = []
    included = []
    prev_sse = sst

    for feature in original_features:
        if feature == "Fuel Type":
            cols = [c for c in X_train.columns if c.startswith("cat__Fuel Type_")]
        else:
            cols = [c for c in X_train.columns if c == f"num__{feature}" or c == feature]
        if not cols:
            continue
        included += cols
        seq_model = sm.OLS(y, sm.add_constant(X_train[included], has_constant="add")).fit()
        curr_sse = float(seq_model.ssr)
        seq_ss = max(0.0, prev_sse - curr_sse)
        reduced_cols = [c for c in X_train.columns if c not in cols]
        reduced_model = sm.OLS(y, sm.add_constant(X_train[reduced_cols], has_constant="add")).fit()
        adj_ss = max(0.0, float(reduced_model.ssr) - full_sse)
        df_term = len(cols)
        adj_ms = adj_ss / df_term
        f_value = adj_ms / mse if mse > 0 else np.nan
        try:
            from scipy.stats import f as f_dist
            p_value = float(f_dist.sf(f_value, df_term, int(model.df_resid)))
        except Exception:
            p_value = np.nan
        rows.append({
            "Sumber": display_name(feature), "DF": df_term, "Seq SS": seq_ss,
            "Kontribusi (%)": (seq_ss / sst * 100) if sst else np.nan,
            "Adj SS": adj_ss, "Adj MS": adj_ms, "F-Value": f_value, "P-Value": p_value
        })
        prev_sse = curr_sse

    rows.extend([
        {"Sumber":"Error", "DF":int(model.df_resid), "Seq SS":full_sse, "Kontribusi (%)":full_sse/sst*100 if sst else np.nan, "Adj SS":full_sse, "Adj MS":mse, "F-Value":np.nan, "P-Value":np.nan},
        {"Sumber":"Total", "DF":int(len(y)-1), "Seq SS":sst, "Kontribusi (%)":100.0, "Adj SS":np.nan, "Adj MS":np.nan, "F-Value":np.nan, "P-Value":np.nan}
    ])
    return pd.DataFrame(rows)


def format_filter_value(value):
    """Readable representation for filter chips and summaries."""
    if pd.isna(value):
        return "Missing"
    if isinstance(value, (float, np.floating)):
        return f"{value:,.2f}".rstrip("0").rstrip(".")
    return str(value)


def apply_buyer_filters(data, filters):
    """Apply user-defined buyer filters to the raw dataset."""
    result = data.copy()
    for col, rule in filters.items():
        if col not in result.columns:
            continue
        kind = rule.get("kind")
        if kind == "numeric":
            lo, hi = rule["range"]
            series = pd.to_numeric(result[col], errors="coerce")
            result = result[series.between(lo, hi, inclusive="both")]
        elif kind == "categorical":
            choices = rule.get("values", [])
            if choices:
                result = result[result[col].astype(str).isin([str(x) for x in choices])]
    return result


def categorical_display_options(col, values):
    vals = list(values)
    if col == "Automatic":
        return [("Manual", 0), ("Automatic", 1)]
    if col == "Metallic":
        return [("Tidak", 0), ("Ya", 1)]
    if col == "Doors":
        return [(str(int(v)), int(v)) for v in sorted(vals, key=lambda x: float(x))]
    if col == "Fuel Type":
        return [(str(v), v) for v in sorted(vals, key=str)]
    return [(str(v), v) for v in sorted(vals, key=str)]


def render_buyer_filter(col, key_prefix, data):
    s = data[col]
    label = display_name(col)
    if col in {"Automatic", "Metallic", "Doors", "Fuel Type"}:
        options = categorical_display_options(col, s.dropna().unique().tolist())
        labels = [x[0] for x in options]
        selected = st.radio(label, ["Semua"] + labels, horizontal=True, key=f"{key_prefix}_radio")
        if selected == "Semua":
            return {"kind": "categorical", "values": []}
        return {"kind": "categorical", "values": [dict(options)[selected]]}

    vals = pd.to_numeric(s, errors="coerce").dropna()
    if vals.empty:
        st.info(f"Tidak ada nilai numerik untuk {label}.")
        return None
    mn, mx = int(vals.min()), int(vals.max())
    chosen = st.slider(label, mn, mx, (mn, mx), key=f"{key_prefix}_range")
    return {"kind": "numeric", "range": chosen}


def simplify_model_name(model):
    """Reduce the 319 raw model strings into useful engine/body type labels."""
    s = str(model).strip()
    s = re.sub(r"(?i)^TOYOTA\s+", "", s)
    s = re.sub(r"(?i)^Corolla\s*", "", s).strip()
    body = None
    for pattern, label in [
        (r"(?i)\bHATCHB\b", "Hatchback"),
        (r"(?i)\bSEDAN\b", "Sedan"),
        (r"(?i)\bLIFTB\b", "Liftback"),
        (r"(?i)\bVERSO\b", "Verso"),
        (r"(?i)\bWAGON\b|\bSTATIONWAGEN\b", "Wagon"),
        (r"(?i)\bMPV\b", "MPV"),
    ]:
        if re.search(pattern, s):
            body = label
            break
    eng16 = re.search(r"(?i)\b(\d(?:\.\d)?)[-\s]?16v\b", s)
    if eng16:
        engine = f"{eng16.group(1)} 16V"
    else:
        eng = re.search(r"(?i)\b(\d(?:\.\d)?)\b", s)
        engine = eng.group(1) if eng else ""
    fuel = "D4D" if re.search(r"(?i)D4D", s) else ""
    parts = []
    if engine:
        parts.append(engine)
    if fuel and fuel not in engine:
        parts.append(fuel)
    if body:
        parts.append(body)
    if not parts:
        return ""
    # Always identify the brand explicitly; all records are Toyota Corolla.
    return "Toyota Corolla " + " ".join(parts)


def chart_interpretation_histogram(data):
    s = pd.to_numeric(data["Price"], errors="coerce").dropna()
    if s.empty:
        return ("Tidak ada data Price yang tersedia.", "Distribusi harga tidak dapat dinilai.", "Periksa kembali data modeling sebelum membaca histogram.")
    skew = s.skew()
    if skew > 0.5:
        shape = "cenderung menceng ke kanan (right-skewed)"
        meaning = "Sebagian besar mobil terkonsentrasi pada harga rendah–menengah, sementara sebagian kecil observasi berada jauh lebih tinggi."
    elif skew < -0.5:
        shape = "cenderung menceng ke kiri (left-skewed)"
        meaning = "Sebagian besar mobil terkonsentrasi pada harga menengah–tinggi, sementara sebagian kecil observasi berada jauh lebih rendah."
    else:
        shape = "relatif seimbang di sekitar pusat distribusi"
        meaning = "Harga tidak menunjukkan ekor yang sangat dominan ke salah satu sisi."
    what = f"Distribusi Price {shape}, dengan mean €{s.mean():,.0f}, median €{s.median():,.0f}, dan skewness {skew:.2f}."
    why = "Bentuk distribusi membantu menentukan bagaimana sebaran Price perlu dibaca dan mengapa observasi pada ekor distribusi perlu diperiksa sebelum pemodelan."
    return what, meaning, why


def chart_interpretation_boxplot(data, col):
    s = pd.to_numeric(data[col], errors="coerce").dropna()
    if s.empty:
        return (f"Tidak ada data {display_name(col)} yang tersedia.", "Boxplot tidak dapat dibentuk.", "Periksa kembali data modeling.")
    q1, med, q3 = s.quantile(.25), s.median(), s.quantile(.75)
    iqr = q3 - q1
    low, high = q1 - 1.5*iqr, q3 + 1.5*iqr
    out = int(((s < low) | (s > high)).sum())
    what = f"Median {display_name(col)} berada di {med:,.0f}, dengan Q1 {q1:,.0f} dan Q3 {q3:,.0f}; IQR = {iqr:,.0f}."
    meaning = f"Sebanyak {out:,} observasi berada di luar batas whisker 1,5×IQR pada data yang sedang aktif."
    why = "Boxplot melengkapi Extreme Value Analysis: IQR digunakan untuk membaca bentuk boxplot, sedangkan treatment model tetap menggunakan batas P1/P99."
    return what, meaning, why


def chart_interpretation_scatter(data, col):
    x = pd.to_numeric(data[col], errors="coerce")
    y = pd.to_numeric(data["Price"], errors="coerce")
    valid = pd.DataFrame({"x": x, "y": y}).dropna()
    if len(valid) < 2 or valid["x"].nunique() < 2:
        return ("Data tidak cukup untuk melihat pola.", "Scatterplot membutuhkan variasi pada variabel yang dipilih.", "Gunakan variabel numerik lain yang memiliki cukup observasi dan variasi.")
    r = float(valid["x"].corr(valid["y"]))
    direction = "positif" if r >= 0 else "negatif"
    strength = "kuat" if abs(r) >= .7 else ("sedang" if abs(r) >= .4 else "lemah")
    return (
        f"Titik data menunjukkan hubungan {direction} dengan Pearson r = {r:.2f}.",
        f"Secara bivariate, kenaikan {display_name(col)} cenderung diikuti {('kenaikan' if r >= 0 else 'penurunan')} Harga Mobil, dengan kekuatan hubungan {strength}.",
        "Garis tren membantu melihat arah hubungan rata-rata, tetapi scatterplot bukan bukti sebab-akibat dan tidak menggantikan multiple regression yang mempertimbangkan predictor secara simultan."
    )


def chart_interpretation_mean_heatmap(data):
    work=data.copy()
    work["Doors"] = pd.to_numeric(work["Doors"], errors="coerce")
    work["Price"] = pd.to_numeric(work["Price"], errors="coerce")
    work=work.dropna(subset=["Fuel Type","Doors","Price"])
    if work.empty:
        return ("Data Fuel Type, Doors, dan Price tidak tersedia.", "Heatmap tidak dapat dibentuk.", "Periksa kembali data modeling.")
    pivot=work.pivot_table(index="Fuel Type", columns="Doors", values="Price", aggfunc="mean")
    if pivot.empty:
        return ("Tidak ada kombinasi Fuel Type × Doors yang tersedia.", "Tidak ada rata-rata kelompok yang dapat dibandingkan.", "Tambahkan data yang memadai sebelum menarik insight kelompok.")
    stacked=pivot.stack()
    idx=stacked.idxmax(); mx=stacked.max(); mn=stacked.min(); idxmin=stacked.idxmin()
    what = f"Rata-rata Price tertinggi terdapat pada {idx[0]} dengan {int(idx[1])} doors, sekitar €{mx:,.0f}; terendah pada {idxmin[0]} dengan {int(idxmin[1])} doors, sekitar €{mn:,.0f}."
    meaning = "Kombinasi Fuel Type dan jumlah Doors menunjukkan perbedaan harga rata-rata antar kelompok dalam dataset."
    why = "Heatmap memberi gambaran awal variasi harga berdasarkan kategori, tetapi perbedaan rata-rata kelompok bukan bukti pengaruh kausal dan tetap perlu dibaca bersama model multivariat."
    return what, meaning, why


def render_insight(parts):
    what, meaning, why = parts
    st.markdown(
        f'<div class="interpretation"><div class="insight-label">WHAT WE SEE</div><div>{html.escape(what)}</div>'
        f'<div class="insight-label">WHAT IT MEANS</div><div>{html.escape(meaning)}</div>'
        f'<div class="insight-label">WHY IT MATTERS</div><div>{html.escape(why)}</div></div>',
        unsafe_allow_html=True,
    )

def build_model_type_table(data, model, transformer, final_feature_names, final_model):
    rows = []
    if "Model" not in data.columns:
        return pd.DataFrame()
    work = data.copy()
    work["Tipe Model"] = work["Model"].map(simplify_model_name)
    for tipe, g in work.groupby("Tipe Model", sort=True):
        base = {}
        for c in ANALYSIS_FEATURES:
            if pd.api.types.is_numeric_dtype(g[c]):
                base[c] = float(pd.to_numeric(g[c], errors="coerce").median())
            else:
                mode = g[c].dropna().mode()
                base[c] = mode.iloc[0] if len(mode) else data[c].dropna().mode().iloc[0]
        inp = pd.DataFrame([base])
        try:
            Xn = transform_new_data(inp[ANALYSIS_FEATURES], transformer).reindex(columns=final_feature_names, fill_value=0.0)
            pred = float(final_model.predict(sm.add_constant(Xn, has_constant="add")).iloc[0])
        except Exception:
            pred = np.nan
        rows.append({"Tipe Model": tipe, "Jumlah Data": len(g), "Median Harga Aktual (€)": g["Price"].median(), "Estimasi Harga Tipikal (€)": pred})
    return pd.DataFrame(rows).sort_values("Jumlah Data", ascending=False).reset_index(drop=True)


st.set_page_config(
    page_title="Corolla — Data Analytics",
    page_icon="○",
    layout="wide",
    initial_sidebar_state="expanded",
)

# ─────────────────────────────────────────────────────────────────────────────
# MINIMAL / EDITORIAL UI
# ─────────────────────────────────────────────────────────────────────────────
st.markdown(
    """
<style>
@import url('https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Manrope:wght@500;600;700;800&display=swap');
:root{--bg:#e9eaeb;--paper:rgba(255,255,255,.58);--white:#fff;--ink:#171717;--muted:#74777b;--soft:#a4a7aa;--line:rgba(0,0,0,.10);--dark:#171717;--gray:#d7d9db;--glass:rgba(255,255,255,.52);--shadow:0 18px 45px rgba(0,0,0,.08),inset 0 1px 0 rgba(255,255,255,.85);--shadow-soft:0 8px 24px rgba(0,0,0,.07),inset 0 1px 0 rgba(255,255,255,.75)}
html,body,[class*="css"]{font-family:'DM Sans',sans-serif;color:var(--ink)!important} body,.stApp{background:var(--bg)!important}.stApp p,.stApp span,.stApp label,.stApp div{color:var(--ink)}.group-badge{display:inline-flex;align-items:center;margin-left:8px;padding:4px 9px;border:1px solid rgba(23,23,23,.16);border-radius:999px;background:rgba(255,255,255,.62);font-size:9px;font-weight:800;letter-spacing:.12em;color:#55595d!important;vertical-align:middle;backdrop-filter:blur(10px);-webkit-backdrop-filter:blur(10px)}
.block-container{max-width:1480px;padding:25px 38px 55px;animation:pageIn .12s ease both}.stApp:before{content:"";position:fixed;inset:0;pointer-events:none;z-index:0;background:radial-gradient(circle at 88% 4%,rgba(255,255,255,.72),transparent 28%),linear-gradient(135deg,#eceeef,#e4e6e8)}.block-container>div{position:relative;z-index:1}
@keyframes pageIn{from{opacity:0;transform:translateY(5px)}to{opacity:1;transform:none}}
@keyframes modalIn{from{opacity:0;transform:translateY(12px) scale(.985)}to{opacity:1;transform:none}}
#MainMenu,footer{visibility:hidden}header{background:transparent!important}[data-testid="stDecoration"]{display:none}
[data-testid="stSidebar"]{background:rgba(232,234,235,.78);backdrop-filter:blur(20px);-webkit-backdrop-filter:blur(20px);border-right:1px solid rgba(0,0,0,.08);box-shadow:10px 0 35px rgba(0,0,0,.05)}[data-testid="stSidebar"]>div:first-child{padding:25px 18px 24px}[data-testid="stSidebar"] *{color:var(--ink)!important}
.sidebar-brand{padding:4px 6px 24px;border-bottom:1px solid var(--line);margin-bottom:17px}.sidebar-brand-mark{display:flex;align-items:center;gap:10px}.toyota-mark{width:34px;height:34px;border:1.5px solid #171717;border-radius:50%;display:flex;align-items:center;justify-content:center;color:#171717;font-weight:800;font-size:13px;background:rgba(255,255,255,.38)}.sidebar-brand-title{font-family:'Manrope',sans-serif;font-size:16px;font-weight:800;letter-spacing:-.03em}.sidebar-brand-sub{color:var(--muted)!important;font-size:10px;margin-top:1px}.nav-caption{font-size:9px;font-weight:800;text-transform:uppercase;letter-spacing:.18em;color:#85898d!important;padding:0 9px 8px}
[data-testid="stSidebar"] [data-testid="stRadio"]>div{gap:4px}[data-testid="stSidebar"] [data-testid="stRadio"] label{padding:9px 11px!important;border-radius:12px;margin:0!important;transition:.16s ease;font-size:12px!important;font-weight:500!important}[data-testid="stSidebar"] [data-testid="stRadio"] label:hover{background:rgba(255,255,255,.52)}[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked){background:#171717;color:#fff!important;font-weight:600!important;box-shadow:0 8px 18px rgba(0,0,0,.14)}[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) p{color:#fff!important}[data-testid="stSidebar"] [data-testid="stRadio"] input{display:none}.sidebar-divider{height:1px;background:var(--line);margin:18px 0}.sidebar-small{color:var(--muted)!important;font-size:10px;line-height:1.55;padding:0 7px}
.topbar{display:flex;align-items:center;justify-content:space-between;margin-bottom:23px}.search-pill{width:310px;height:38px;border:1px solid rgba(0,0,0,.08);background:rgba(255,255,255,.42);backdrop-filter:blur(14px);border-radius:22px;display:flex;align-items:center;padding:0 15px;color:#8b8f93;font-size:11px;box-shadow:var(--shadow-soft)}.search-icon{font-size:16px;margin-right:9px;color:#444}.top-date{color:#666b70!important;font-size:11px}.avatar{display:inline-flex;align-items:center;justify-content:center;width:34px;height:34px;border-radius:50%;background:#171717;color:#fff!important;margin-left:12px;font-size:10px;font-weight:700}
.eyebrow{color:#777b7f!important;font-size:9px;text-transform:uppercase;letter-spacing:.20em;font-weight:800;margin-bottom:7px}h1,h2,h3{font-family:'Manrope',sans-serif;color:var(--ink)}.hero-title{color:#111!important;font-family:'Manrope',sans-serif;font-size:46px;line-height:1.02;letter-spacing:-.055em;font-weight:800;margin:0}.hero-copy{color:#656a6f!important;font-size:13px;line-height:1.6;max-width:650px;margin-top:12px}.page-title{color:#111!important;font-family:'Manrope',sans-serif;font-size:33px;line-height:1.08;letter-spacing:-.045em;font-weight:800;margin:0}.page-copy{color:#686d72!important;font-size:12px;margin:8px 0 0}
.hero-wrap{min-height:260px;padding:34px 38px 30px;border:1px solid rgba(255,255,255,.65);border-radius:28px;background:linear-gradient(135deg,rgba(255,255,255,.62),rgba(255,255,255,.30));backdrop-filter:blur(18px);position:relative;overflow:hidden;margin-bottom:18px;box-shadow:var(--shadow)}
.stButton button,.stDownloadButton button,.stFormSubmitButton button{border-radius:11px!important;min-height:39px;font-size:11px;font-weight:700;border:1px solid rgba(0,0,0,.10)!important;background:rgba(255,255,255,.52)!important;color:#171717!important;box-shadow:var(--shadow-soft)!important;transition:transform .16s ease,background .16s ease,box-shadow .16s ease}.stButton button:hover,.stDownloadButton button:hover,.stFormSubmitButton button:hover{border-color:#171717!important;background:#171717!important;color:#fff!important;transform:translateY(-1px);box-shadow:0 10px 22px rgba(0,0,0,.15)!important}.primary-btn button{background:#171717!important;color:#fff!important;border-color:#171717!important}
.kpi-strip{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin-bottom:18px}.kpi-cell{padding:17px 19px;min-height:92px;border:1px solid rgba(255,255,255,.68);border-radius:17px;background:rgba(255,255,255,.47);backdrop-filter:blur(16px);box-shadow:var(--shadow-soft)}.kpi-label{color:#74797e!important;font-size:9px;text-transform:uppercase;letter-spacing:.10em;font-weight:800}.kpi-value{font-family:'Manrope',sans-serif;font-size:25px;font-weight:800;letter-spacing:-.04em;margin-top:7px;color:#151515!important}.kpi-note{color:#858a8f!important;font-size:9px;margin-top:2px}
.card{background:rgba(255,255,255,.48);border:1px solid rgba(255,255,255,.70);border-radius:20px;padding:20px;box-shadow:var(--shadow-soft);backdrop-filter:blur(16px);-webkit-backdrop-filter:blur(16px)}.card-dark{background:#171717;color:white;border-color:#171717}.card-title{font-family:'Manrope',sans-serif;font-size:15px;font-weight:700;letter-spacing:-.025em;color:#151515!important}.card-sub{color:#777c81!important;font-size:10px;margin-top:4px}.muted{color:var(--muted)!important}.rule{height:1px;background:var(--line);margin:14px 0}.quick{display:flex;align-items:center;justify-content:space-between;padding:11px 0;border-bottom:1px solid rgba(0,0,0,.08);font-size:11px}.quick:last-child{border-bottom:0}
.result-card{background:#171717;color:#fff;border-radius:22px;padding:27px;margin-top:17px;position:relative;overflow:hidden;box-shadow:0 18px 35px rgba(0,0,0,.16)}.result-card:after{content:"";position:absolute;width:230px;height:230px;border:1px solid rgba(255,255,255,.12);border-radius:50%;right:-75px;top:-80px;box-shadow:0 0 0 35px rgba(255,255,255,.035)}.result-label{color:#aaa!important;text-transform:uppercase;letter-spacing:.16em;font-size:9px;font-weight:700}.result-value{font-family:'Manrope',sans-serif;font-size:42px;font-weight:800;letter-spacing:-.055em;margin:6px 0 3px;color:#fff!important}.result-meta{color:#bdbdb7!important;font-size:10px}.section-head{display:flex;align-items:end;justify-content:space-between;margin:22px 0 10px}.section-head-title{font-family:'Manrope',sans-serif;font-size:15px;font-weight:800;letter-spacing:-.035em;color:#151515!important}.section-head-sub{color:#777c81!important;font-size:10px}.soft-chip{display:inline-flex;align-items:center;gap:6px;padding:6px 10px;border-radius:999px;background:rgba(255,255,255,.45);border:1px solid rgba(0,0,0,.08);font-size:9px;color:#666b70!important}
.stTextInput input,.stNumberInput input,[data-baseweb="select"] *{color:#202020!important;background:rgba(255,255,255,.68)!important}.stSelectbox>div>div,.stMultiSelect>div>div{border-radius:11px!important;background:rgba(255,255,255,.58)!important;border-color:rgba(0,0,0,.09)!important}.stSlider [data-baseweb="slider"] div{color:#171717}.stCheckbox label,.stRadio label{font-size:11px!important}.stRadio>div{gap:6px}
[data-testid="stDataFrame"]{border:1px solid rgba(0,0,0,.08);border-radius:15px;overflow:hidden;box-shadow:0 8px 20px rgba(0,0,0,.05);background:rgba(255,255,255,.42)}[data-testid="stDataFrame"] *{color:#202020}.stTabs [data-baseweb="tab-list"]{gap:5px;background:rgba(0,0,0,.055);padding:5px;border-radius:15px;border:1px solid rgba(255,255,255,.55);backdrop-filter:blur(12px)}.stTabs [data-baseweb="tab"]{height:36px;padding:0 17px;border-radius:10px;color:#666b70!important;font-size:11px;font-weight:600}.stTabs [aria-selected="true"]{background:#171717!important;color:#fff!important;box-shadow:0 7px 16px rgba(0,0,0,.12)}.stTabs [data-baseweb="tab-highlight"]{background:transparent!important}
/* table/header neutral palette */
[data-testid="stDataFrame"] thead th{background:#252627!important;color:#fff!important}.dataframe thead th{background:#252627!important;color:#fff!important}.dataframe tbody td{background:rgba(255,255,255,.55)!important}
.editorial-dark .stButton button{background:rgba(255,255,255,.10)!important;color:#fff!important;border-color:rgba(255,255,255,.16)!important;box-shadow:none!important}.editorial-dark .stButton button:hover{background:#fff!important;color:#171717!important}.editorial-dark .rule{background:rgba(255,255,255,.14)}
/* No decorative black grid blocks: keep the editorial layout clean and content-first. */
.editorial-dark{background:#4b4e52;color:#fff;border-radius:20px;padding:22px;box-shadow:0 16px 30px rgba(0,0,0,.12)}.editorial-dark *{color:#fff!important}
.add-filter-wrap [data-testid="stButton"] button{background:#171717!important;color:#fff!important;border-color:#171717!important;box-shadow:0 10px 22px rgba(0,0,0,.13)!important}.add-filter-wrap [data-testid="stButton"] button:hover{background:#303234!important;color:#fff!important}
/* modal + page transitions */
[data-testid="stDialog"]>div{animation:modalIn .22s cubic-bezier(.2,.8,.2,1) both;background:rgba(245,246,247,.90)!important;backdrop-filter:blur(22px);border:1px solid rgba(255,255,255,.8);box-shadow:0 24px 70px rgba(0,0,0,.20)}[data-testid="stDialog"] [data-testid="stDialogHeader"]{background:rgba(255,255,255,.35)!important}

/* Predict Price: pre-populated criteria panel, matching Find a Car's visual language. */
.prediction-criteria-card .section-head{margin-top:4px}.prediction-criteria-card [data-testid="stRadio"]{padding:5px 0 10px}.prediction-criteria-card [data-testid="stRadio"]>div{gap:7px;flex-wrap:wrap}.prediction-criteria-card [data-testid="stRadio"] label{border-radius:999px!important;padding:4px 9px!important;background:transparent!important;border:1px solid transparent!important;transition:.16s ease}.prediction-criteria-card [data-testid="stRadio"] label:hover{background:rgba(255,255,255,.48)!important}.prediction-criteria-card [data-testid="stSelectbox"]>div>div{min-height:40px}.prediction-note{margin-top:12px;padding:12px 15px;border:1px solid rgba(0,0,0,.08);border-radius:14px;background:rgba(255,255,255,.46);backdrop-filter:blur(12px);color:#686d72!important;font-size:10px;line-height:1.55}
@media(max-width:900px){.block-container{padding:18px}.kpi-strip{grid-template-columns:1fr 1fr}.hero-title{font-size:36px}.search-pill{width:210px}}

/* Premium monochrome controls: ONLY white / gray / charcoal. Never red. */
:root{--primary-color:#777b7f!important;--secondary-background-color:#e7e8e9!important}
.stApp{--primary-color:#777b7f!important}
input[type="radio"],input[type="checkbox"]{accent-color:#777b7f!important}
/* Streamlit/BaseWeb selected indicators */
[data-testid="stRadio"] [role="radio"][aria-checked="true"],
[data-testid="stCheckbox"] [role="checkbox"][aria-checked="true"],
[data-baseweb="radio"] [aria-checked="true"]{background:#777b7f!important;border-color:#777b7f!important;color:#fff!important}
[data-testid="stRadio"] [role="radio"][aria-checked="true"] svg,
[data-testid="stCheckbox"] [role="checkbox"][aria-checked="true"] svg,
[data-baseweb="radio"] [aria-checked="true"] svg{color:#fff!important;fill:#fff!important;stroke:#fff!important}
[data-testid="stCheckbox"] input:checked + div{background:#777b7f!important;border-color:#777b7f!important}
[data-baseweb="slider"] [role="slider"]{background:#777b7f!important;border-color:#777b7f!important}
[data-baseweb="slider"] [data-testid="stSliderTrackFill"]{background:#777b7f!important}
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked){background:#777b7f!important;color:#fff!important;font-weight:600!important}
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) *{color:#fff!important}
[data-testid="stSidebar"] [data-testid="stRadio"] label:hover{background:rgba(119,123,127,.12)!important}
/* Streamlit widgets sometimes expose accent through pseudo-elements. Force the neutral gray. */
[data-testid="stRadio"] label div[role="radio"], [data-testid="stCheckbox"] label div[role="checkbox"]{color:#777b7f!important}
[data-testid="stRadio"] label div[role="radio"][aria-checked="true"], [data-testid="stCheckbox"] label div[role="checkbox"][aria-checked="true"]{background:#777b7f!important;border-color:#777b7f!important;color:#fff!important}
/* Clickable tabs: inactive = white/black; active = charcoal/white. */
.stTabs [data-baseweb="tab"]{background:rgba(255,255,255,.72)!important;color:#171717!important;border:1px solid rgba(0,0,0,.09)!important}
.stTabs [data-baseweb="tab"] *{color:#171717!important}
.stTabs [data-baseweb="tab"][aria-selected="true"]{background:#171717!important;color:#ffffff!important;border-color:#171717!important;border-bottom:0!important;box-shadow:0 7px 16px rgba(0,0,0,.12)!important}
.stTabs [data-baseweb="tab"][aria-selected="true"] *{color:#ffffff!important;background:transparent!important}
.stTabs [data-baseweb="tab-highlight"]{display:none!important;background:transparent!important;height:0!important}
.stTabs [data-baseweb="tab-border"]{display:none!important}
/* Active subpage: dark background with gray text ONLY when selected. No red indicator. */
.stTabs [data-baseweb="tab"][aria-selected="true"], .stTabs [data-baseweb="tab"][aria-selected="true"] *{color:#ffffff!important}
.stTabs [data-baseweb="tab"][aria-selected="true"]{border-bottom:0!important;outline:none!important}
.stTabs [data-baseweb="tab"][aria-selected="true"]::after{display:none!important;background:transparent!important}
.stTabs [data-baseweb="tab-highlight"], .stTabs [data-baseweb="tab-border"]{display:none!important;background:transparent!important;border:0!important}


/* Final active-tab contrast: selected tab = black background + white text only. */
.stTabs [data-baseweb="tab"][aria-selected="true"] span,
.stTabs [data-baseweb="tab"][aria-selected="true"] div,
.stTabs [data-baseweb="tab"][aria-selected="true"] p{color:#ffffff!important;}
.stTabs [data-baseweb="tab"][aria-selected="true"]::before,
.stTabs [data-baseweb="tab"][aria-selected="true"]::after{background:transparent!important;border:0!important;}

/* Dark cards are always high-contrast: never dark text on a dark background. */
.editorial-dark,.editorial-dark *,.result-card,.result-card *{color:#fff!important}.result-card{background:#4b4e52!important}.card-dark{background:#4b4e52!important;color:#fff!important}.card-dark *{color:#fff!important}

.interpretation{margin-top:14px;padding:13px 15px;border:1px solid rgba(0,0,0,.08);border-radius:13px;background:rgba(255,255,255,.52);color:#5f6265!important;font-size:11px;line-height:1.65}
.interpretation b{color:#303234!important}
[data-testid="stDialog"]{animation:modalIn .16s cubic-bezier(.2,.8,.2,1) both}

/* FINAL TAB OVERRIDE: active tab stays white with black text. No red indicator. */
.stTabs [data-baseweb="tab"] { 
    background: #ffffff !important;
    color: #171717 !important;
    border: 1px solid rgba(0,0,0,.10) !important;
    border-bottom: 0 !important;
    box-shadow: none !important;
}
.stTabs [data-baseweb="tab"] * { color: #171717 !important; background: transparent !important; }
.stTabs [data-baseweb="tab"][aria-selected="true"] {
    background: #ffffff !important;
    color: #171717 !important;
    border-color: rgba(0,0,0,.14) !important;
    border-bottom: 0 !important;
    box-shadow: none !important;
}
.stTabs [data-baseweb="tab"][aria-selected="true"] * {
    color: #171717 !important;
    background: transparent !important;
}
.stTabs [data-baseweb="tab-highlight"],
.stTabs [data-baseweb="tab-border"],
.stTabs [data-baseweb="tab"][aria-selected="true"]::before,
.stTabs [data-baseweb="tab"][aria-selected="true"]::after {
    display: none !important;
    background: transparent !important;
    border: 0 !important;
}

/* V17 FINAL: neutral Streamlit controls and tabs. Active tab is WHITE + BLACK TEXT. */
/* Remove every red primary/accent treatment from native controls. */
:root{
  --primary-color:#777b7f !important;
  --primary-color-rgb:119,123,127 !important;
  --secondary-background-color:#e7e8e9 !important;
}
/* Sidebar radio: selected pill = gray with white text; unselected = light with black text. */
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked){
  background:#777b7f !important;
  color:#ffffff !important;
  box-shadow:0 8px 18px rgba(0,0,0,.10) !important;
}
[data-testid="stSidebar"] [data-testid="stRadio"] label:has(input:checked) *{
  color:#ffffff !important;
}
[data-testid="stSidebar"] [data-testid="stRadio"] label:not(:has(input:checked)){
  background:transparent !important;
  color:#171717 !important;
}
[data-testid="stSidebar"] [data-testid="stRadio"] label:not(:has(input:checked)) *{
  color:#171717 !important;
}
/* Native radio/check controls: gray only. */
[data-testid="stRadio"] [role="radio"][aria-checked="true"],
[data-testid="stRadio"] div[role="radio"][aria-checked="true"],
[data-baseweb="radio"] [role="radio"][aria-checked="true"]{
  background-color:#777b7f !important;
  border-color:#777b7f !important;
  color:#ffffff !important;
}
[data-testid="stRadio"] [role="radio"][aria-checked="true"] svg,
[data-baseweb="radio"] [role="radio"][aria-checked="true"] svg{
  color:#ffffff !important;
  fill:#ffffff !important;
  stroke:#ffffff !important;
}
[data-testid="stCheckbox"] [role="checkbox"][aria-checked="true"],
[data-baseweb="checkbox"] [role="checkbox"][aria-checked="true"]{
  background-color:#777b7f !important;
  border-color:#777b7f !important;
}
/* Slider: gray filled track and thumb, never red. */
[data-testid="stSlider"] [role="slider"]{
  background:#777b7f !important;
  border-color:#777b7f !important;
}
[data-testid="stSlider"] [data-baseweb="slider"] div[style*="background"]{
  background:#777b7f !important;
}
[data-testid="stSlider"] [data-baseweb="slider"] > div > div:first-child{
  background:#777b7f !important;
}
/* FINAL TAB BEHAVIOR: inactive = transparent, active = white. */
[data-testid="stTabs"] [role="tab"],
[data-testid="stTabs"] [data-baseweb="tab"]{
  background:transparent !important;
  background-color:transparent !important;
  color:#171717 !important;
  border:1px solid transparent !important;
  border-bottom:0 !important;
  box-shadow:none !important;
}
[data-testid="stTabs"] [role="tab"] *,
[data-testid="stTabs"] [data-baseweb="tab"] *{
  color:#171717 !important;
  background:transparent !important;
  -webkit-text-fill-color:#171717 !important;
}
[data-testid="stTabs"] [role="tab"][aria-selected="true"],
[data-testid="stTabs"] [data-baseweb="tab"][aria-selected="true"]{
  background:#ffffff !important;
  background-color:#ffffff !important;
  color:#171717 !important;
  border:1px solid rgba(0,0,0,.10) !important;
  border-bottom:0 !important;
  box-shadow:0 5px 14px rgba(0,0,0,.045) !important;
}
[data-testid="stTabs"] [role="tab"][aria-selected="true"] *,
[data-testid="stTabs"] [data-baseweb="tab"][aria-selected="true"] *{
  color:#171717 !important;
  background:transparent !important;
  -webkit-text-fill-color:#171717 !important;
}
/* Remove Streamlit/BaseWeb selected indicators, especially the default red line. */
[data-testid="stTabs"] [data-baseweb="tab-highlight"],
[data-testid="stTabs"] [data-baseweb="tab-border"],
[data-testid="stTabs"] [role="tab"]::before,
[data-testid="stTabs"] [role="tab"]::after,
[data-testid="stTabs"] [data-baseweb="tab"]::before,
[data-testid="stTabs"] [data-baseweb="tab"]::after{
  background:transparent !important;
  background-color:transparent !important;
  border:0 !important;
  box-shadow:none !important;
  display:none !important;
}
[data-testid="stTabs"] [data-baseweb="tab-list"] > div[style*="background"],
[data-testid="stTabs"] [data-baseweb="tab-list"] > div[style*="border-bottom"]{
  background:transparent !important;
  border-bottom-color:transparent !important;
}


/* V19 visual system: neutral gray only, no red accents. */
.top-meta{display:flex;align-items:center;justify-content:flex-end;gap:12px;margin-bottom:20px;color:#666b70!important;font-size:11px}.top-meta *{color:#666b70!important}.top-meta .avatar{color:#fff!important;background:#171717!important}
.search-pill{display:none!important}
[data-testid="stAppViewContainer"] button, [data-testid="stSidebar"] button{accent-color:#777b7f!important}
.stTabs [data-baseweb="tab"]{background:transparent!important;color:#171717!important;border:1px solid transparent!important;box-shadow:none!important}
.stTabs [data-baseweb="tab"] *{color:#171717!important;-webkit-text-fill-color:#171717!important}
.stTabs [data-baseweb="tab"][aria-selected="true"]{background:#ffffff!important;color:#171717!important;border:1px solid rgba(0,0,0,.08)!important;box-shadow:0 8px 20px rgba(0,0,0,.06)!important}
.stTabs [data-baseweb="tab"][aria-selected="true"] *{color:#171717!important;-webkit-text-fill-color:#171717!important}
.stTabs [data-baseweb="tab-highlight"],.stTabs [data-baseweb="tab-border"]{display:none!important;background:transparent!important}
.equation-note{margin-top:14px;padding:14px 16px;border-radius:14px;background:rgba(238,242,245,.78);border:1px solid rgba(0,0,0,.06);color:#4f555a!important;font-size:11px;line-height:1.65}

/* V20: active tabs use a white glass surface; inactive tabs are transparent. */
.stTabs [data-baseweb="tab"], .stTabs [role="tab"]{background:transparent!important;background-color:transparent!important;color:#171717!important;border:1px solid transparent!important;border-radius:14px!important;box-shadow:none!important;backdrop-filter:none!important;-webkit-backdrop-filter:none!important}
.stTabs [data-baseweb="tab"] *, .stTabs [role="tab"] *{color:#171717!important;-webkit-text-fill-color:#171717!important;background:transparent!important}
.stTabs [data-baseweb="tab"][aria-selected="true"], .stTabs [role="tab"][aria-selected="true"]{background:rgba(255,255,255,.72)!important;background-color:rgba(255,255,255,.72)!important;color:#171717!important;border:1px solid rgba(255,255,255,.95)!important;border-radius:14px!important;box-shadow:0 8px 24px rgba(31,35,38,.08),inset 0 1px 0 rgba(255,255,255,.98)!important;backdrop-filter:blur(16px) saturate(120%)!important;-webkit-backdrop-filter:blur(16px) saturate(120%)!important}
.stTabs [data-baseweb="tab"][aria-selected="true"] *, .stTabs [role="tab"][aria-selected="true"] *{color:#171717!important;-webkit-text-fill-color:#171717!important;background:transparent!important}
/* V19.1 final neutral controls: all native selection accents are gray. */
input[type="radio"], input[type="checkbox"]{accent-color:#777b7f!important;}
[data-baseweb="radio"] > div:first-child,[data-baseweb="radio"] div[role="radio"],
[data-testid="stRadio"] div[role="radio"]{border-color:#b8bdc1!important;}
[data-baseweb="radio"] > div:first-child[aria-checked="true"],
[data-baseweb="radio"] div[role="radio"][aria-checked="true"],
[data-testid="stRadio"] div[role="radio"][aria-checked="true"]{background:#777b7f!important;border-color:#777b7f!important;}
[data-baseweb="radio"] > div:first-child[aria-checked="true"]::after,
[data-baseweb="radio"] div[role="radio"][aria-checked="true"]::after{background:#ffffff!important;border-color:#ffffff!important;}
[data-testid="stSlider"] [data-baseweb="slider"] div{background-color:#777b7f!important;}
[data-testid="stSlider"] [role="slider"]{background:#777b7f!important;border-color:#777b7f!important;}
[data-testid="stCheckbox"] [role="checkbox"]{border-color:#b8bdc1!important;}
[data-testid="stCheckbox"] [role="checkbox"][aria-checked="true"]{background:#777b7f!important;border-color:#777b7f!important;}
/* Final tab interaction: inactive tabs are text-only; active tab gets a white rounded surface. */
.stTabs [data-baseweb="tab-list"]{background:transparent!important;border:0!important;padding:3px!important;gap:6px!important}
.stTabs [data-baseweb="tab"]{background:transparent!important;background-color:transparent!important;border:1px solid transparent!important;border-radius:12px!important;box-shadow:none!important;color:#171717!important;padding:0 16px!important}
.stTabs [data-baseweb="tab"] *{background:transparent!important;color:#171717!important;-webkit-text-fill-color:#171717!important}
.stTabs [data-baseweb="tab"][aria-selected="true"]{background:#fff!important;background-color:#fff!important;border:1px solid rgba(0,0,0,.08)!important;border-radius:12px!important;box-shadow:0 7px 18px rgba(0,0,0,.06)!important}
.stTabs [data-baseweb="tab"][aria-selected="true"] *{background:transparent!important;color:#171717!important;-webkit-text-fill-color:#171717!important}
.stTabs [data-baseweb="tab-highlight"],.stTabs [data-baseweb="tab-border"]{display:none!important;background:transparent!important;border:0!important}
</style>
""",
    unsafe_allow_html=True,
)

# ─────────────────────────────────────────────────────────────────────────────
# DATA
# ─────────────────────────────────────────────────────────────────────────────
# Dataset is bundled with the app; no upload step is required.
df_raw = load_dataset(None)

if "Price" not in df_raw.columns:
    st.error("Kolom `Harga Mobil` tidak ditemukan. Pastikan dataset memiliki kolom target Harga Mobil.")
    st.stop()

all_cols = df_raw.columns.tolist()
eligible = [c for c in ANALYSIS_FEATURES if c in df_raw.columns]
default_features = eligible.copy()

# Extreme-value analysis is always calculated from the untouched/raw dataset.
extreme_summary, extreme_observations, extreme_mask = extreme_value_analysis(df_raw)

# ─────────────────────────────────────────────────────────────────────────────
# SIDEBAR
# ─────────────────────────────────────────────────────────────────────────────
st.sidebar.markdown(
    """
<div class="sidebar-brand">
  <div class="sidebar-brand-mark">
    <div class="toyota-mark">T</div>
    <div>
      <div class="sidebar-brand-title">Toyota Corolla</div>
      <div class="sidebar-brand-sub">Data Analytics</div>
    </div>
  </div>
</div>
<div class="nav-caption">Workspace</div>
""",
    unsafe_allow_html=True,
)

page_options = ["Dashboard", "Find & Predict Car Price"]
if "nav_selection" not in st.session_state:
    st.session_state.nav_selection = "Dashboard"
if "pending_page" in st.session_state:
    pending = st.session_state.pop("pending_page")
    if pending in page_options:
        st.session_state.nav_selection = pending

def navigate_to(target):
    # Do not mutate the radio widget's own state during its active run.
    st.session_state.pending_page = target

page = st.sidebar.radio(
    "Navigation",
    page_options,
    key="nav_selection",
    label_visibility="collapsed",
)

st.sidebar.markdown('<div class="sidebar-divider"></div>', unsafe_allow_html=True)
st.sidebar.markdown('<div class="nav-caption">Model Utama</div>', unsafe_allow_html=True)
feature_cols = eligible.copy()
st.sidebar.caption("Model utama menggunakan 11 variabel: 1 target + 10 predictor. Id dan Model tidak digunakan sebagai predictor.")
split_pct = st.sidebar.slider("Data pelatihan (%)", 60, 90, 80, 5)
standardize = st.sidebar.checkbox("Standarisasi variabel numerik", False)
outlier_treatment = st.sidebar.radio(
    "Penanganan extreme value",
    ["Keep Outliers", "Remove Outliers"],
    index=1,
)
compare_both_models = st.sidebar.checkbox("Compare Both Models", value=False)

st.sidebar.markdown('<div class="sidebar-divider"></div>', unsafe_allow_html=True)
st.sidebar.markdown(
    f'<div class="sidebar-small"><b>{len(df_raw):,}</b> observasi dimuat<br>'
    f'<b>{len(df_raw.columns)}</b> kolom · target <b>Harga Mobil</b><br><br>'
    f'<b>{len(eligible)+1}</b> variabel digunakan · <b>Id</b> dan <b>Model</b> tidak digunakan sebagai predictor.<br>'
    f'Outlier treatment: <b>{html.escape(outlier_treatment)}</b>.<br>'
    'Tipe model dipakai sebagai filter pencarian, bukan predictor regresi.</div>', 
    unsafe_allow_html=True,
)

if not feature_cols:
    st.warning("Pilih minimal satu variabel prediktor di sidebar.")
    st.stop()

# ─────────────────────────────────────────────────────────────────────────────
# MODEL
# ─────────────────────────────────────────────────────────────────────────────
FINAL_FEATURES = eligible.copy()

# Raw/original data is never modified. Only the modeling copy changes when the
# user selects Remove Outliers.
if outlier_treatment == "Remove Outliers":
    df_model = df_raw.loc[~extreme_mask].copy()
else:
    df_model = df_raw.copy()


@st.cache_data(show_spinner=False)
def prepare_final_model(df_json, features, target, split, scale):
    df = pd.read_json(io.StringIO(df_json))
    bundle = prepare_regression_split(df, list(features), target, test_size=1 - split / 100, standardize=scale, seed=42)
    mt = compute_metrics(bundle["y_train"], bundle["pred_train"], len(bundle["feature_names"]))
    mv = compute_metrics(bundle["y_test"], bundle["pred_test"], len(bundle["feature_names"]))
    coef = extract_coefficients(bundle["model"], bundle["feature_names"])
    return bundle, mt, mv, coef


with st.spinner("Preparing model..."):
    final_bundle, final_metrics_train, final_metrics_test, final_coef_df = prepare_final_model(
        df_model.to_json(), tuple(FINAL_FEATURES), "Price", split_pct, standardize
    )

final_model = final_bundle["model"]
final_transformer = final_bundle["transformer"]
final_full_model, final_metrics_full = fit_full_model_metrics(final_bundle)
final_feature_names = final_bundle["feature_names"]
final_prep_report = final_bundle["report"]
final_y_test = final_bundle["y_test"]
final_y_pred_test = final_bundle["pred_test"]
naive_prediction = np.repeat(float(final_bundle["y_train"].mean()), len(final_y_test))
naive_metrics = compute_metrics(final_y_test, naive_prediction, 0)

# Backward-compatible aliases used by existing dashboard/export sections.

metrics_train = final_metrics_train
metrics_test = final_metrics_test
prep_report = final_prep_report
coef_df = final_coef_df
y_test = final_y_test
y_pred_test = final_y_pred_test

# Keep a separate comparison only when explicitly requested. Both models use
# the same predictor set, split percentage, random seed, and standardization.
@st.cache_data(show_spinner=False)
def compare_models(df_json, clean_json, features, split, scale):
    raw_df = pd.read_json(io.StringIO(df_json))
    clean_df = pd.read_json(io.StringIO(clean_json))
    bundle_a, mt_a, mv_a, _ = prepare_final_model(df_json, features, "Price", split, scale)
    bundle_b, mt_b, mv_b, _ = prepare_final_model(clean_json, features, "Price", split, scale)
    table = pd.DataFrame([
        {
            "Model": "Model A — Keep Outliers",
            "Observasi": len(raw_df),
            "R² Train": mt_a["r2"],
            "R² Test": mv_a["r2"],
            "Adjusted R²": mv_a["adj_r2"],
            "MAE Test": mv_a["mae"],
            "RMSE Test": mv_a["rmse"],
            "MAPE Test": mv_a["mape"],
        },
        {
            "Model": "Model B — Remove Outliers",
            "Observasi": len(clean_df),
            "R² Train": mt_b["r2"],
            "R² Test": mv_b["r2"],
            "Adjusted R²": mv_b["adj_r2"],
            "MAE Test": mv_b["mae"],
            "RMSE Test": mv_b["rmse"],
            "MAPE Test": mv_b["mape"],
        },
    ])
    return table

comparison_table = None
if compare_both_models:
    comparison_table = compare_models(
        df_raw.to_json(), df_model.to_json(), tuple(FINAL_FEATURES), split_pct, standardize
    )


def plot_layout(fig, height=None):
    fig.update_layout(
        template="plotly_white",
        paper_bgcolor="rgba(0,0,0,0)",
        plot_bgcolor="rgba(0,0,0,0)",
        font=dict(family="DM Sans", color="#5f605b", size=11),
        title=None,
        margin=dict(l=10, r=10, t=10, b=10),
        legend=dict(bgcolor="rgba(0,0,0,0)", font=dict(size=10)),
        height=height or 300,
    )
    fig.update_xaxes(showgrid=True, gridcolor="#ecece8", zeroline=False, linecolor="#d7d7d1")
    fig.update_yaxes(showgrid=True, gridcolor="#ecece8", zeroline=False, linecolor="#d7d7d1")
    return fig



def show_chart(fig, title, key):
    """Render a clean Plotly chart with a meaningful chart title and toolbar."""
    fig.update_layout(
        title=dict(
            text=title,
            x=0.02,
            xanchor="left",
            font=dict(family="DM Sans", size=15, color="#171717"),
        ),
        margin=dict(l=10, r=10, t=48, b=10),
    )
    st.plotly_chart(
        fig,
        use_container_width=True,
        config={
            "displayModeBar": True,
            "displaylogo": False,
            "scrollZoom": True,
            "responsive": True,
            "toImageButtonOptions": {
                "format": "png",
                "filename": key,
                "scale": 2,
            },
        },
        key=f"plot_{key}",
    )

def style_table(df):
    """Consistent Swiss/editorial dark-gray table header styling."""
    return df.style.set_table_styles([
        {"selector": "th", "props": [("background-color", "#252627"), ("color", "#ffffff"), ("font-weight", "700"), ("border-color", "#252627")]},
        {"selector": "td", "props": [("border-color", "#e1e2e3")]},
    ])


def build_full_analysis_xlsx(df_raw, df_model, feature_cols, metrics_train, metrics_test, coef_df, y_test, y_pred_test, extreme_summary, comparison_table):
    predictions = pd.DataFrame({"Harga Aktual": np.asarray(y_test), "Harga Prediksi": np.asarray(y_pred_test), "Residual": np.asarray(y_test)-np.asarray(y_pred_test)})
    coef_export = coef_df.copy(); coef_export["Feature"] = coef_export["Feature"].map(display_name)
    metrics_df = pd.DataFrame({"Metric":["R²","Adjusted R²","MAE","RMSE","MAPE"],"Training":[metrics_train["r2"],metrics_train["adj_r2"],metrics_train["mae"],metrics_train["rmse"],metrics_train["mape"]],"Test":[metrics_test["r2"],metrics_test["adj_r2"],metrics_test["mae"],metrics_test["rmse"],metrics_test["mape"]]})
    nums = df_model[[c for c in ["Price"]+list(feature_cols) if c in df_model.columns]].select_dtypes(include=np.number)
    rows=[]
    for c in nums.columns:
        ss=pd.to_numeric(nums[c],errors="coerce").dropna(); rows.append({"Variabel":display_name(c),"Count":int(ss.count()),"Mean":ss.mean(),"Std":ss.std(),"Min":ss.min(),"Q1":ss.quantile(.25),"Median":ss.median(),"Q3":ss.quantile(.75),"Max":ss.max()})
    xlsx=io.BytesIO()
    with pd.ExcelWriter(xlsx,engine="openpyxl") as writer:
        df_raw.rename(columns=DISPLAY_LABELS).to_excel(writer,sheet_name="Dataset",index=False)
        pd.DataFrame(rows).to_excel(writer,sheet_name="EDA Result",index=False)
        metrics_df.to_excel(writer,sheet_name="Model Metrics",index=False)
        coef_export.to_excel(writer,sheet_name="Coefficients",index=False)
        predictions.to_excel(writer,sheet_name="Predictions",index=False)
        extreme_summary.to_excel(writer,sheet_name="Extreme Values",index=False)
        if comparison_table is not None: comparison_table.to_excel(writer,sheet_name="Model Comparison",index=False)
        from openpyxl.styles import PatternFill
        for ws in writer.book.worksheets:
            ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
            for cell in ws[1]:
                cell.font=cell.font.copy(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor="33373A")
            for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(max(len(str(c.value)) if c.value is not None else 0 for c in col)+2,42)
    return xlsx.getvalue()

def kpi_strip(items):
    cells = []
    for label, value, note in items:
        cells.append(
            f'<div class="kpi-cell"><div class="kpi-label">{html.escape(label)}</div>'
            f'<div class="kpi-value">{html.escape(str(value))}</div>'
            f'<div class="kpi-note">{html.escape(str(note))}</div></div>'
        )
    st.markdown('<div class="kpi-strip">' + ''.join(cells) + '</div>', unsafe_allow_html=True)


def section_head(title, subtitle=""):
    st.markdown(
        f'<div class="section-head"><div><div class="section-head-title">{html.escape(title)}</div>'
        f'<div class="section-head-sub">{html.escape(subtitle)}</div></div></div>',
        unsafe_allow_html=True,
    )


def predict_input(input_df):
    X_new = transform_new_data(input_df[FINAL_FEATURES], final_transformer)
    X_new = X_new.reindex(columns=final_feature_names, fill_value=0.0)
    pred = float(final_model.predict(sm.add_constant(X_new, has_constant="add")).iloc[0])
    return max(pred, 0)


def _interval_options(series, width, unit_label="", start_at_zero=False):
    """Create human-readable dropdown intervals from the observed dataset range."""
    vals = pd.to_numeric(series, errors="coerce").dropna()
    if vals.empty:
        return []
    lo = int(np.floor(vals.min()))
    hi = int(np.ceil(vals.max()))
    start = 0 if start_at_zero else max(0, (lo // width) * width)
    options = []
    current = start
    while current <= hi:
        end = current + width
        label = f"{current:,}–{end:,}{(' ' + unit_label) if unit_label else ''}"
        mask = vals.between(current, end, inclusive="left")
        observed = vals[mask]
        representative = float(observed.median()) if not observed.empty else float((current + end) / 2)
        options.append((label, representative))
        current = end
    return options


def _model_scoped_values(selected_model, feature):
    """Return values actually observed for the selected simplified model type."""
    work = df_raw.copy()
    if selected_model and "Model" in work.columns:
        work = work[work["Model"].map(simplify_model_name) == selected_model]
    vals = pd.to_numeric(work[feature], errors="coerce").dropna().unique().tolist()
    return sorted(vals)


def form_values(prefix, selected_model):
    """Build prediction inputs as guided dropdowns, scoped to the selected model type."""
    vals = {}
    cols = st.columns(2)
    scoped = df_raw.copy()
    if selected_model and "Model" in scoped.columns:
        scoped = scoped[scoped["Model"].map(simplify_model_name) == selected_model]
    if scoped.empty:
        scoped = df_raw.copy()

    # Weight dan Quart Tax tidak ditampilkan sebagai kriteria pencarian harga.
    # Jika keduanya tetap dibutuhkan oleh model final, nilainya diisi otomatis
    # menggunakan median observasi pada tipe model yang dipilih.
    input_features = [f for f in FINAL_FEATURES if f not in {"Weight", "Quart Tax"}]
    for i, feat in enumerate(input_features):
        col = cols[i % 2]
        s = scoped[feat]

        if feat == "Age":
            options = _interval_options(s, 12, "bulan", start_at_zero=True)
            labels = [x[0] for x in options]
            chosen = col.selectbox(display_name(feat), labels, key=f"{prefix}_{feat}", help="Pilih interval usia. Model menggunakan median observasi pada interval tersebut sebagai nilai numerik representatif.")
            vals[feat] = dict(options)[chosen]

        elif feat == "Kilometers":
            options = _interval_options(s, 25000, "km", start_at_zero=True)
            labels = [x[0] for x in options]
            chosen = col.selectbox(display_name(feat), labels, key=f"{prefix}_{feat}", help="Pilih interval jarak tempuh. Nilai yang dikirim ke model adalah median observasi pada interval tersebut.")
            vals[feat] = dict(options)[chosen]

        elif feat == "CC":
            vals_num = pd.to_numeric(s, errors="coerce").dropna()
            choices = sorted(vals_num.unique().tolist())
            if not choices:
                choices = sorted(pd.to_numeric(df_raw[feat], errors="coerce").dropna().unique().tolist())
            formatted = [f"{int(v):,}" if float(v).is_integer() else f"{v:,.2f}" for v in choices]
            chosen = col.selectbox(display_name(feat), formatted, key=f"{prefix}_{feat}", help="Pilihan dibatasi pada nilai yang benar-benar tersedia untuk tipe model yang dipilih.")
            vals[feat] = choices[formatted.index(chosen)]

        elif feat in {"Automatic", "Metallic", "Doors"}:
            options = categorical_display_options(feat, s.dropna().unique().tolist())
            labels = [x[0] for x in options]
            chosen = col.selectbox(display_name(feat), labels, key=f"{prefix}_{feat}")
            vals[feat] = dict(options)[chosen]

        elif feat == "Fuel Type":
            options = categorical_display_options(feat, s.dropna().unique().tolist())
            labels = [x[0] for x in options]
            chosen = col.selectbox(display_name(feat), labels, key=f"{prefix}_{feat}")
            vals[feat] = dict(options)[chosen]

        else:
            choices = sorted(s.dropna().astype(str).unique().tolist())
            if choices:
                chosen = col.selectbox(display_name(feat), choices, key=f"{prefix}_{feat}")
                vals[feat] = chosen
            else:
                vals[feat] = np.nan

# Predictor yang tidak dipakai sebagai kriteria UI tetap diisi secara
    # otomatis agar struktur input persis sesuai dengan model regresi final.
    for hidden_feat in ("Weight", "Quart Tax"):
        if hidden_feat in FINAL_FEATURES:
            hidden_values = pd.to_numeric(scoped[hidden_feat], errors="coerce").dropna()
            if hidden_values.empty:
                hidden_values = pd.to_numeric(df_raw[hidden_feat], errors="coerce").dropna()
            vals[hidden_feat] = float(hidden_values.median()) if not hidden_values.empty else 0.0

    return pd.DataFrame([vals])

# ─────────────────────────────────────────────────────────────────────────────
# PAGE RENDERERS
# ─────────────────────────────────────────────────────────────────────────────
def render_find_car():
    st.markdown('<div class="eyebrow">Buyer workspace</div><div class="page-title">Find the Corolla that fits your needs.</div><div class="page-copy">Choose a model, then add only the variables you care about. The search updates to your own criteria.</div>', unsafe_allow_html=True)
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)

    # Model is a first-class filter. Streamlit's selectbox is searchable by typing.
    model_labels = sorted([x for x in df_raw["Model"].dropna().map(simplify_model_name).unique().tolist() if x]) if "Model" in df_raw.columns else []
    model_options = ["Semua tipe"] + model_labels
    selected_model = st.selectbox("Tipe Model", model_options, index=0, help="Ketik tipe Corolla yang ingin dicari.")

    if "buyer_filters" not in st.session_state:
        st.session_state.buyer_filters = {}

    available_filter_cols = [c for c in eligible if c != "Model" and c not in st.session_state.buyer_filters]
    st.markdown('<div class="card">', unsafe_allow_html=True)
    section_head("Custom filters", "Tambah variabel dari variabel yang bisa digunakan untuk pencarian")
    add_col = st.selectbox(
        "Tambah variabel",
        available_filter_cols,
        format_func=display_name,
        key="buyer_add_filter_col",
        placeholder="Pilih variabel yang ingin dicari...",
    ) if available_filter_cols else None
    if add_col:
        c1, c2 = st.columns([1, 4])
        with c1:
            st.markdown('<div class="add-filter-wrap">', unsafe_allow_html=True)
            if st.button("+ Tambah", use_container_width=True, key="add_buyer_filter"):
                st.session_state.buyer_filters[add_col] = {"kind": "pending"}
                st.rerun()
            st.markdown('</div>', unsafe_allow_html=True)
        with c2:
            st.caption(f"Variabel tersedia: {len(eligible)} · filter aktif: {len(st.session_state.buyer_filters)}")
    else:
        st.caption(f"{len(eligible)} variabel tersedia untuk pencarian. Filter aktif: {len(st.session_state.buyer_filters)}")
    st.markdown('</div>', unsafe_allow_html=True)

    # Render active filters and update their values.
    active_filters = {}
    if st.session_state.buyer_filters:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Your criteria", "Atur nilai setiap variabel sesuai kebutuhan")
        for idx, col in enumerate(list(st.session_state.buyer_filters.keys())):
            top1, top2 = st.columns([5, 1])
            with top1:
                rule = render_buyer_filter(col, f"buyer_{idx}_{col}", df_raw)
                if rule is not None:
                    active_filters[col] = rule
            with top2:
                st.markdown('<div style="height:26px"></div>', unsafe_allow_html=True)
                if st.button("Hapus", key=f"remove_buyer_{col}", use_container_width=True):
                    del st.session_state.buyer_filters[col]
                    st.rerun()
        st.session_state.buyer_filters = active_filters
        st.markdown('</div>', unsafe_allow_html=True)

    # Apply model + custom criteria.
    search_filters = dict(active_filters)
    filtered = df_raw.copy()
    if selected_model != "Semua tipe" and "Model" in filtered.columns:
        filtered = filtered[filtered["Model"].map(simplify_model_name) == selected_model]
    filtered = apply_buyer_filters(filtered, search_filters)

    st.markdown('<div class="card">', unsafe_allow_html=True)
    section_head("Matching cars", "Hasil yang memenuhi model dan semua kriteria pilihanmu")
    k1, k2, k3 = st.columns(3)
    k1.metric("Mobil cocok", f"{len(filtered):,}")
    k2.metric("Tipe dipilih", "Semua" if selected_model == "Semua tipe" else "1 tipe")
    k3.metric("Filter aktif", f"{len(active_filters)}")

    if len(filtered) == 0:
        st.warning("Tidak ada mobil yang cocok. Coba longgarkan salah satu filter.")
    else:
        show_cols = [c for c in ["Model", "Price"] + ANALYSIS_FEATURES if c in filtered.columns and c != "Id"]
        extra_cols = [c for c in active_filters if c not in show_cols and c in filtered.columns]
        table_cols = show_cols + extra_cols
        display_df = filtered[table_cols].rename(columns=DISPLAY_LABELS).copy()
        st.dataframe(style_table(display_df), use_container_width=True, hide_index=True, height=500)
        csv = filtered.rename(columns=DISPLAY_LABELS).to_csv(index=False).encode("utf-8")
        st.download_button("Export hasil pencarian", csv, "hasil_pencarian_toyota_corolla.csv", "text/csv")
    st.markdown('</div>', unsafe_allow_html=True)


def render_predict_price():
    st.markdown('<div class="eyebrow">Prediction</div><div class="page-title">Estimate a used-car price.</div><div class="page-copy">Atur seluruh spesifikasi kendaraan langsung seperti panel <b>Your criteria</b> pada Find a Car. Semua variabel model tersedia di sini, termasuk Pajak Tahunan dan Bobot Kendaraan.</div>', unsafe_allow_html=True)
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)

    pred_model_options = sorted([x for x in df_raw["Model"].dropna().map(simplify_model_name).unique().tolist() if x]) if "Model" in df_raw.columns else []
    selected_pred_model = st.selectbox(
        "Tipe Model", pred_model_options, key="single_pred_model",
        help="Pilih tipe Toyota Corolla sebagai konteks kendaraan. Model tidak digunakan sebagai predictor regresi."
    ) if pred_model_options else None

    st.markdown('<div class="card prediction-criteria-card">', unsafe_allow_html=True)
    section_head("Your criteria", "Atur nilai setiap variabel sesuai kendaraan yang ingin diprediksi")
    st.caption("Jenis bahan bakar mengikuti seluruh kategori yang tersedia pada dataset: CNG, Diesel, dan Petrol.")

    scoped = df_raw.copy()
    if selected_pred_model and "Model" in scoped.columns:
        scoped = scoped[scoped["Model"].map(simplify_model_name) == selected_pred_model]
    if scoped.empty:
        scoped = df_raw.copy()

    vals = {}
    with st.form("prediction_criteria_form"):
        # Numeric predictors use the same range-slider language as Find a Car.
        # For prediction, the midpoint of the selected range is passed to the model.
        numeric_features = [f for f in FINAL_FEATURES if f not in {"Automatic", "Metallic", "Doors", "Fuel Type"}]
        categorical_features = [f for f in FINAL_FEATURES if f in {"Automatic", "Metallic", "Doors", "Fuel Type"}]
        ordered_features = []
        # Keep the same visual two-column rhythm as Find a Car.
        for f in FINAL_FEATURES:
            if f not in ordered_features:
                ordered_features.append(f)

        for row_start in range(0, len(ordered_features), 2):
            row_features = ordered_features[row_start:row_start + 2]
            cols = st.columns(2, gap="large")
            for col_ui, feat in zip(cols, row_features):
                s = scoped[feat]
                with col_ui:
                    if feat in {"Automatic", "Metallic", "Doors", "Fuel Type"}:
                        # Fuel Type harus selalu menampilkan seluruh kategori yang memang
                        # ada di dataset (termasuk CNG), walaupun tipe model yang dipilih
                        # kebetulan tidak memiliki observasi CNG.
                        option_source = s.dropna().unique().tolist()
                        if feat == "Fuel Type":
                            option_source = df_raw[feat].dropna().unique().tolist()
                        options = categorical_display_options(feat, option_source)
                        labels = [x[0] for x in options]
                        chosen = st.radio(
                            display_name(feat), labels, horizontal=True,
                            key=f"pred_criteria_{feat}"
                        )
                        vals[feat] = dict(options)[chosen]
                    else:
                        vals_num = pd.to_numeric(s, errors="coerce").dropna()
                        if vals_num.empty:
                            vals_num = pd.to_numeric(df_raw[feat], errors="coerce").dropna()
                        if vals_num.empty:
                            st.info(f"Tidak ada nilai untuk {display_name(feat)}.")
                            vals[feat] = np.nan
                            continue

                        mn = int(np.floor(vals_num.min()))
                        mx = int(np.ceil(vals_num.max()))
                        if mn == mx:
                            chosen_range = (mn, mx)
                        else:
                            # Start at the observed median, represented as a one-point range.
                            med = int(round(float(vals_num.median())))
                            med = min(max(med, mn), mx)
                            chosen_range = st.slider(
                                display_name(feat), mn, mx, (med, med),
                                key=f"pred_criteria_{feat}_range",
                                help="Geser batas kiri/kanan. Untuk prediksi, nilai tengah dari rentang yang dipilih digunakan sebagai nilai input model."
                            )
                        if mn == mx:
                            st.slider(display_name(feat), mn, mx, (mn, mx), key=f"pred_criteria_{feat}_range_fixed")
                        vals[feat] = float(sum(chosen_range) / 2)

        submit = st.form_submit_button("Hitung Estimasi Harga  →", use_container_width=True)

    st.markdown('</div>', unsafe_allow_html=True)
    st.markdown('<div class="prediction-note">Semua 10 predictor model aktif di panel ini. Slider numerik mengikuti pola <b>Your criteria</b> pada Find a Car, sedangkan variabel kategorikal dipilih langsung. Nilai tengah rentang numerik digunakan sebagai input prediksi.</div>', unsafe_allow_html=True)

    if submit:
        inp = pd.DataFrame([vals])
        pred = predict_input(inp)
        margin = 1.96 * final_metrics_test["rmse"]
        lo, hi = max(0, pred - margin), pred + margin
        st.markdown(
            f'<div class="result-card"><div class="result-label">Estimated price</div>'
            f'<div class="result-value">€{pred:,.0f}</div>'
            f'<div class="result-meta">Approximate model error band: €{lo:,.0f} — €{hi:,.0f}</div></div>',
            unsafe_allow_html=True,
        )

        # Cari kendaraan yang benar-benar memenuhi kriteria yang baru dipilih.
        # Model type mengikuti simplified model, numerik mengikuti rentang slider,
        # dan kategorikal harus sama persis dengan pilihan pengguna.
        matching = df_raw.copy()
        if selected_pred_model and "Model" in matching.columns:
            matching = matching[matching["Model"].map(simplify_model_name) == selected_pred_model]

        # Simpan rentang numerik dari widget untuk digunakan lagi pada matching.
        numeric_ranges = {}
        for feat in FINAL_FEATURES:
            if feat in {"Automatic", "Metallic", "Doors", "Fuel Type"}:
                continue
            key = f"pred_criteria_{feat}_range"
            if key in st.session_state:
                numeric_ranges[feat] = st.session_state[key]

        for feat, (left, right) in numeric_ranges.items():
            if feat in matching.columns:
                vals_series = pd.to_numeric(matching[feat], errors="coerce")
                matching = matching[vals_series.between(left, right, inclusive="both")]

        for feat in {"Automatic", "Metallic", "Doors", "Fuel Type"}.intersection(FINAL_FEATURES):
            if feat in vals and feat in matching.columns:
                try:
                    matching = matching[matching[feat] == vals[feat]]
                except Exception:
                    pass

        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Matching Cars", "Kendaraan yang memenuhi seluruh kriteria yang dipilih pada panel prediksi")

        if matching.empty:
            st.info("Belum ada kendaraan yang benar-benar cocok dengan kombinasi kriteria ini. Coba lebarkan salah satu rentang numerik atau ubah pilihan kategorikal.")
        else:
            # Statistik hanya dihitung dari mobil yang match, bukan seluruh dataset.
            median_actual = float(pd.to_numeric(matching["Price"], errors="coerce").median())
            try:
                # Prediksi seluruh mobil yang match untuk mendapatkan harga tipikal model.
                X_match = transform_new_data(matching[FINAL_FEATURES], final_transformer)
                X_match = X_match.reindex(columns=final_feature_names, fill_value=0.0)
                pred_match = np.asarray(final_model.predict(sm.add_constant(X_match, has_constant="add")), dtype=float)
                typical_pred = float(np.median(pred_match))
            except Exception:
                typical_pred = float(pred)

            m1, m2, m3 = st.columns(3)
            m1.metric("Mobil cocok", f"{len(matching):,}")
            m2.metric("Median harga aktual", f"€{median_actual:,.0f}")
            m3.metric("Harga tipikal model", f"€{typical_pred:,.0f}")

            show_cols = [c for c in ["Model", "Price"] + ANALYSIS_FEATURES if c in matching.columns and c != "Id"]
            display_matching = matching[show_cols].rename(columns=DISPLAY_LABELS).copy()
            st.dataframe(style_table(display_matching), use_container_width=True, hide_index=True, height=420)

        st.markdown('</div>', unsafe_allow_html=True)


def render_model_evaluation():
    st.markdown('<div class="eyebrow">Regression model</div><div class="page-title">Evaluate how well the model performs.</div><div class="page-copy">Model final untuk prediksi menggunakan 10 predictor dari materi; sidebar dapat digunakan untuk eksperimen model alternatif.</div>', unsafe_allow_html=True)
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)

    st.markdown('<div class="card">', unsafe_allow_html=True)
    section_head("Dataset yang digunakan", "Ringkasan dataset Excel terbaru dan konfigurasi model saat ini")
    d1, d2, d3, d4 = st.columns(4)
    d1.metric("Observasi", f"{len(df_raw):,}")
    d2.metric("Kolom", f"{len(df_raw.columns)}")
    d3.metric("Predictor aktif", f"{len(feature_cols)}")
    d4.metric("Fitur setelah encoding", f"{prep_report['n_features_out']}")
    st.markdown('<div style="height:10px"></div>', unsafe_allow_html=True)
    active_table = pd.DataFrame({
        "Variabel aktif": [display_name(c) for c in feature_cols],
        "Tipe data": [str(df_raw[c].dtype) for c in feature_cols],
    })
    st.dataframe(style_table(active_table), use_container_width=True, hide_index=True, height=260)
    st.markdown('</div>', unsafe_allow_html=True)

    kpi_strip([
        ("R²", f"{final_metrics_test['r2']:.4f}", "test goodness of fit"),
        ("MAE", f"€{final_metrics_test['mae']:,.0f}", "mean absolute error"),
        ("RMSE", f"€{final_metrics_test['rmse']:,.0f}", "root mean squared error"),
        ("MAPE", f"{final_metrics_test['mape']:.2f}%", "mean absolute percentage error"),
    ])

    st.markdown('<div class="card">', unsafe_allow_html=True)
    section_head("Naive Benchmark vs Regression", "Baseline menggunakan rata-rata harga training pada seluruh data test")
    benchmark_df = pd.DataFrame({
        "Model": ["Naive benchmark", "Multiple Linear Regression"],
        "R²": [naive_metrics["r2"], final_metrics_test["r2"]],
        "MAE": [naive_metrics["mae"], final_metrics_test["mae"]],
        "RMSE": [naive_metrics["rmse"], final_metrics_test["rmse"]],
        "MAPE": [naive_metrics["mape"], final_metrics_test["mape"]],
    })
    st.dataframe(benchmark_df.style.set_table_styles([{ "selector": "th", "props": [("background-color", "#252627"), ("color", "#ffffff"), ("font-weight", "700")] }]).format({"R²":"{:.4f}","MAE":"€ {:,.0f}","RMSE":"€ {:,.0f}","MAPE":"{:.2f}%"}), use_container_width=True, hide_index=True)
    st.markdown('</div>', unsafe_allow_html=True)

    st.markdown('<div class="card">', unsafe_allow_html=True)
    section_head("Outlier Treatment", "Perbandingan data asli dan data setelah menghapus observasi extreme P1/P99")
    ot1, ot2, ot3 = st.columns(3)
    ot1.metric("Data awal", f"{len(df_raw):,}")
    ot2.metric("Data extreme", f"{int(extreme_mask.sum()):,}")
    ot3.metric("Data setelah removal", f"{len(df_model):,}")
    st.caption("Raw/original dataset tidak diubah. Removal hanya diterapkan pada salinan data yang digunakan untuk modeling.")
    if compare_both_models and comparison_table is not None:
        st.markdown('<div style="height:8px"></div>', unsafe_allow_html=True)
        st.markdown('<div class="card-title">Model Comparison</div>', unsafe_allow_html=True)
        st.dataframe(comparison_table.style.set_table_styles([{ "selector": "th", "props": [("background-color", "#252627"), ("color", "#ffffff"), ("font-weight", "700")] }]).format({
            "R² Train":"{:.4f}", "R² Test":"{:.4f}", "Adjusted R²":"{:.4f}",
            "MAE Test":"€ {:,.0f}", "RMSE Test":"€ {:,.0f}", "MAPE Test":"{:.2f}%"
        }), use_container_width=True, hide_index=True)
        keep_row = comparison_table.iloc[0]
        remove_row = comparison_table.iloc[1]
        r2_better = "Keep Outliers" if keep_row["R² Test"] >= remove_row["R² Test"] else "Remove Outliers"
        mae_better = "Keep Outliers" if keep_row["MAE Test"] <= remove_row["MAE Test"] else "Remove Outliers"
        rmse_better = "Keep Outliers" if keep_row["RMSE Test"] <= remove_row["RMSE Test"] else "Remove Outliers"
        mape_better = "Keep Outliers" if keep_row["MAPE Test"] <= remove_row["MAPE Test"] else "Remove Outliers"
        st.markdown(
            f"<div class='soft-chip'>Test R² terbaik: <b>{r2_better}</b></div> &nbsp; "
            f"<div class='soft-chip'>MAE terbaik: <b>{mae_better}</b></div> &nbsp; "
            f"<div class='soft-chip'>RMSE terbaik: <b>{rmse_better}</b></div> &nbsp; "
            f"<div class='soft-chip'>MAPE terbaik: <b>{mape_better}</b></div>",
            unsafe_allow_html=True,
        )
        st.caption("Perbandingan dinilai dari data test. Remove Outliers tidak otomatis lebih baik; gunakan performa test dan konteks analisis untuk menentukan model final.")
    else:
        st.caption("Aktifkan 'Compare Both Models' di sidebar untuk membandingkan Keep Outliers dan Remove Outliers.")
    st.markdown('</div>', unsafe_allow_html=True)

    a, b = st.columns([1.15, .85])
    with a:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Harga Aktual dan Harga Prediksi", "Membandingkan hasil prediksi pada data test")
        fig = px.scatter(pd.DataFrame({"Harga Aktual": final_y_test.values, "Harga Prediksi": final_y_pred_test.values}), x="Harga Aktual", y="Harga Prediksi", opacity=.62)
        fig.update_traces(marker=dict(size=5, color="#222222"))
        lo = min(float(final_y_test.min()), float(final_y_pred_test.min())); hi = max(float(final_y_test.max()), float(final_y_pred_test.max()))
        fig.add_trace(go.Scatter(x=[lo, hi], y=[lo, hi], mode="lines", line=dict(color="#666666", width=2), name="Ideal"))
        fig = plot_layout(fig, 350)
        fig.update_xaxes(title="Harga Aktual (€)"); fig.update_yaxes(title="Harga Prediksi (€)")
        show_chart(fig, "Harga Aktual vs Harga Prediksi", "model_actual_pred")
        st.markdown('</div>', unsafe_allow_html=True)
    with b:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Perbandingan Performa Train dan Test", f"Pembagian data: {split_pct}% train ({len(final_bundle['y_train']):,} observasi) dan {100-split_pct}% test ({len(final_bundle['y_test']):,} observasi)")
        comp = pd.DataFrame({
            "Data": ["Train", "Test"],
            "R²": [final_metrics_train["r2"], final_metrics_test["r2"]],
        })
        fig = px.bar(comp, x="Data", y="R²", text="R²")
        fig.update_traces(marker_color="#222222", texttemplate="%{text:.3f}", textposition="outside")
        fig.update_yaxes(range=[0, 1], tickformat=".0%", title="R²")
        fig.update_xaxes(title="Pembagian Data")
        fig = plot_layout(fig, 350)
        show_chart(fig, "Perbandingan R² Train dan Test", "model_train_test")
        st.caption(f"MAE / RMSE / MAPE tetap tersedia pada tabel Model Evaluation. Nilai-nilai tersebut akan dihitung ulang ketika pembagian train-test digeser dari {split_pct}% / {100-split_pct}%.")
        st.markdown('</div>', unsafe_allow_html=True)



# TOP META
st.markdown(
    f"""<div class="top-meta"><span>Regression workspace · {len(FINAL_FEATURES)+1} variables used</span><span class="avatar">DA</span></div>""",
    unsafe_allow_html=True,
)

# ─────────────────────────────────────────────────────────────────────────────
# DASHBOARD
# ─────────────────────────────────────────────────────────────────────────────
if page == "Dashboard":
    st.markdown('<div class="page-title">Toyota Corolla Price Analytics</div><div class="page-copy">Analisis harga mobil Toyota Corolla bekas menggunakan Multiple Linear Regression. <span class="group-badge">KELOMPOK 5 · KELAS A</span></div>', unsafe_allow_html=True)
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)
    kpi_strip([
        ("Jumlah Observasi", f"{len(df_model):,}", f"setelah {outlier_treatment.lower()}"),
        ("Jumlah Variabel", f"{len(FINAL_FEATURES)+1}", "1 target + 10 predictor"),
        ("Target", "Harga Mobil", "dalam Euro (€)"),
        ("Model", "Multiple Linear Regression", "OLS · 10 predictor"),
    ])

    dash_tabs = st.tabs(["Dataset & Model", "Histogram", "Boxplot", "Heatmap", "Scatterplot", "Residuals", "Extreme Values"])

    with dash_tabs[0]:
        # Regression equation comes first, followed by EDA, model evaluation, then dataset.
        st.markdown('<div class="card" style="background:rgba(255,255,255,.66);">', unsafe_allow_html=True)
        section_head("Regression Equation", "Persamaan model OLS yang digunakan untuk memprediksi Harga Mobil")
        st.markdown('<div class="equation-note"><b>Rumus umum Multiple Linear Regression</b><br><span style="font-family:Georgia,serif;font-size:19px">ŷ = β₀ + β₁X₁ + β₂X₂ + ··· + βₖXₖ</span><br>β₀ adalah intercept dan β₁…βₖ adalah koefisien predictor. OLS memilih koefisien yang meminimalkan jumlah kuadrat residual pada data training.</div>', unsafe_allow_html=True)
        intercept=float(final_coef_df.loc[final_coef_df["Feature"]=="Intercept","Coefficient"].iloc[0])
        expr=rf"\text{{Harga Mobil}} = {intercept:,.2f}"
        for _,row in final_coef_df[final_coef_df["Feature"]!="Intercept"].iterrows():
            coef=float(row["Coefficient"]); sign="+" if coef>=0 else "-"
            safe=str(display_name(row["Feature"])).replace("_",r"\_").replace(" ",r"\ ")
            expr += rf" {sign} {abs(coef):,.2f}\times\text{{{safe}}}"
        st.latex(expr)
        st.markdown('<div class="equation-note"><b>Dari mana angka koefisien?</b><br>Semua angka pada persamaan merupakan hasil estimasi <b>OLS dari training data</b> setelah preprocessing dan treatment extreme value yang sedang aktif. Koefisien dibaca sebagai perubahan rata-rata prediksi Harga Mobil untuk kenaikan 1 unit predictor, dengan predictor lain dianggap konstan. Untuk Fuel Type, koefisien dummy dibaca terhadap reference category.</div>', unsafe_allow_html=True)
        coef_explain=final_coef_df.copy(); coef_explain["Variabel"]=coef_explain["Feature"].map(display_name); coef_explain["Koefisien"]=coef_explain["Coefficient"]
        coef_explain["Penjelasan"] = coef_explain.apply(lambda r: "Intercept/β₀: konstanta hasil estimasi OLS" if str(r["Feature"])=="Intercept" else f"Koefisien OLS untuk {display_name(r['Feature'])}: perubahan rata-rata Harga Mobil untuk kenaikan 1 unit, predictor lain konstan", axis=1)
        st.dataframe(style_table(coef_explain[["Variabel","Koefisien","Penjelasan"]]).format({"Koefisien":"{:,.2f}"}),use_container_width=True,hide_index=True)
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="card" style="margin-top:16px">', unsafe_allow_html=True)
        section_head("Koefisien Model OLS", "Rincian koefisien, uji t, p-value, interval kepercayaan 95%, dan VIF")
        vif_detail = build_vif_table(final_model)
        coef_detail = final_coef_df.copy()
        coef_detail["Term"] = coef_detail["Feature"].map(display_name)
        coef_detail = coef_detail.merge(vif_detail, on="Feature", how="left")
        coef_detail = coef_detail[["Term","Coefficient","t-Statistic","p-value","CI Lower (95%)","CI Upper (95%)","VIF"]]
        coef_detail.columns = ["Term","Koefisien","T-Value","P-Value","CI 95% Lower","CI 95% Upper","VIF"]
        st.dataframe(style_table(coef_detail).format({"Koefisien":"{:,.4f}","T-Value":"{:,.2f}","P-Value":"{:.4f}","CI 95% Lower":"{:,.4f}","CI 95% Upper":"{:,.4f}","VIF":"{:,.2f}"}), use_container_width=True, hide_index=True, height=min(560,90+len(coef_detail)*38))
        st.caption("Fuel Type diperlakukan sebagai variabel kategorikal dengan dummy coding. Satu kategori menjadi reference category; VIF dihitung pada fitur hasil encoding yang masuk ke OLS.")

        section_head("Ringkasan Model", "Ringkasan R² training, test, dan full model untuk membedakan fit dan evaluasi prediksi")
        model_summary = build_model_summary_table(final_model, final_metrics_train, final_metrics_test, final_metrics_full)
        st.dataframe(style_table(model_summary).format({"Nilai":"{:,.4f}"}), use_container_width=True, hide_index=True, height=355)
        st.caption("R² Training berasal dari data yang digunakan untuk fitting model. R² Test berasal dari data yang disisihkan untuk evaluasi. R² Full Model dihitung dengan mem-fit ulang OLS pada seluruh observasi setelah treatment aktif, sehingga dapat digunakan sebagai pembanding dengan output software yang mem-fit seluruh data.")

        section_head("Analysis of Variance (ANOVA)", "Kontribusi dan signifikansi predictor berdasarkan training data")
        anova_table = build_anova_table(final_model, final_bundle["X_train"], final_bundle["y_train"], FINAL_FEATURES)
        st.dataframe(style_table(anova_table).format({"Seq SS":"{:,.0f}","Kontribusi (%)":"{:.2f}","Adj SS":"{:,.0f}","Adj MS":"{:,.0f}","F-Value":"{:,.2f}","P-Value":lambda x: "—" if pd.isna(x) else f"{x:.4f}"}), use_container_width=True, hide_index=True, height=min(560,100+len(anova_table)*38))
        st.caption("Seq SS menunjukkan kontribusi berdasarkan urutan predictor pada tabel. Adj SS menunjukkan kontribusi setelah predictor lain diperhitungkan. Hasil ANOVA berasal dari training data yang digunakan untuk fitting OLS.")
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="card" style="margin-top:16px">', unsafe_allow_html=True)
        nums_for_eda = df_model[[c for c in ["Price"] + FINAL_FEATURES if c in df_model.columns]].select_dtypes(include=np.number)
        eda_rows = []
        for c in nums_for_eda.columns:
            s_num = pd.to_numeric(nums_for_eda[c], errors="coerce").dropna()
            eda_rows.append({"Variabel": display_name(c), "Count": int(s_num.count()), "Mean": s_num.mean(), "Std": s_num.std(), "Min": s_num.min(), "Q1": s_num.quantile(.25), "Median": s_num.quantile(.50), "Q3": s_num.quantile(.75), "Max": s_num.max()})
        eda_table = pd.DataFrame(eda_rows)
        section_head("EDA Result · Statistik Deskriptif", "Ringkasan numerik setelah treatment data yang sedang aktif")
        st.dataframe(style_table(eda_table).format({c:"{:,.2f}" for c in ["Mean","Std","Min","Q1","Median","Q3","Max"]}), use_container_width=True, hide_index=True, height=min(430, 72+len(eda_table)*36))
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="card" style="margin-top:16px">', unsafe_allow_html=True)
        section_head("Model Evaluation", "Performa model pada data test · satu metric per baris")
        eval_table = pd.DataFrame({
            "Metric": ["R²", "Adjusted R²", "MAE", "RMSE", "MAPE"],
            "Test Result": [final_metrics_test["r2"], final_metrics_test["adj_r2"], final_metrics_test["mae"], final_metrics_test["rmse"], final_metrics_test["mape"]],
        })
        st.dataframe(style_table(eval_table).format({"Test Result": lambda x: f"{x:.3f}"}), use_container_width=True, hide_index=True, height=235)
        st.markdown(f'<div class="interpretation"><div class="insight-label">WHAT WE SEE</div><div>Test R² = <b>{final_metrics_test["r2"]:.3f}</b>, Adjusted R² = <b>{final_metrics_test["adj_r2"]:.3f}</b>, MAE = <b>€{final_metrics_test["mae"]:,.0f}</b>, RMSE = <b>€{final_metrics_test["rmse"]:,.0f}</b>, dan MAPE = <b>{final_metrics_test["mape"]:.2f}%</b>.</div><div class="insight-label">WHAT IT MEANS</div><div>Model menjelaskan sekitar <b>{final_metrics_test["r2"]*100:.1f}%</b> variasi Price pada data test; rata-rata error absolut sekitar <b>€{final_metrics_test["mae"]:,.0f}</b>.</div><div class="insight-label">WHY IT MATTERS</div><div>Evaluasi test menunjukkan kemampuan generalisasi model ke observasi yang tidak digunakan untuk fitting. MAE, RMSE, dan MAPE melengkapi R² karena langsung menggambarkan besar kesalahan prediksi.</div></div>', unsafe_allow_html=True)
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div class="card" style="margin-top:16px">', unsafe_allow_html=True)
        section_head("Dataset", "Seluruh observasi dapat discroll; tidak dibatasi 100 baris")
        st.dataframe(style_table(df_raw.rename(columns=DISPLAY_LABELS)), use_container_width=True, hide_index=True, height=620)
        xlsx_data=build_full_analysis_xlsx(df_raw,df_model,FINAL_FEATURES,final_metrics_train,final_metrics_test,final_coef_df,final_y_test,final_y_pred_test,extreme_summary,comparison_table)
        export_spacer, export_col = st.columns([5.5, 1.5])
        with export_col:
            st.download_button("Export Hasil Analisis ↓",xlsx_data,"Toyota_Corolla_Full_Analysis.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",use_container_width=True,key="dashboard_full_analysis_export")
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[1]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Distribusi Harga Mobil Toyota Corolla", "Melihat bentuk dan konsentrasi sebaran Harga Mobil")
        # Minitab-style histogram: equal-width bins with explicit boundaries.
        # The old 35-bin setting was arbitrary; using 30 equal-width intervals
        # gives a cleaner distribution view while keeping every observation in
        # its correct interval. Minitab also uses equally spaced bins and
        # frequency on the y-axis by default.
        price = pd.to_numeric(df_model["Price"], errors="coerce").dropna().to_numpy()
        if len(price):
            hist_counts, hist_edges = np.histogram(price, bins=30)
            hist_centers = (hist_edges[:-1] + hist_edges[1:]) / 2
            bin_width = hist_edges[1] - hist_edges[0]
            hist_fig = go.Figure(
                go.Bar(
                    x=hist_centers,
                    y=hist_counts,
                    width=bin_width * 0.98,
                    marker_color="#171717",
                    marker_line_color="#171717",
                    customdata=np.column_stack([hist_edges[:-1], hist_edges[1:]]),
                    hovertemplate=(
                        "Interval: €%{customdata[0]:,.0f} – €%{customdata[1]:,.0f}"
                        "<br>Frequency: %{y}<extra></extra>"
                    ),
                )
            )
            hist_fig.update_xaxes(title="Harga Mobil (€)")
            hist_fig.update_yaxes(title="Frequency")
            fig = hist_fig
        else:
            fig = go.Figure()
        fig.update_traces(marker_line_width=1)
        show_chart(plot_layout(fig,460), "Distribusi Harga Mobil Toyota Corolla", "dash_hist")
        render_insight(chart_interpretation_histogram(df_model))
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[2]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Sebaran Variabel dan Harga Mobil", "Melihat median, kuartil, dan observasi ekstrem")
        chosen=st.selectbox("Variabel", [c for c in EXTREME_NUMERIC_FEATURES if c in df_raw.columns], format_func=display_name, key="dash_box_var")
        fig=px.box(df_model,y=chosen,points="outliers"); fig.update_traces(marker_color="#777777",line_color="#171717")
        show_chart(plot_layout(fig,460), f"Sebaran {display_name(chosen)} pada Harga Mobil", "dash_box")
        render_insight(chart_interpretation_boxplot(df_model, chosen))
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[3]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Korelasi Antarvariabel Numerik", "Melihat arah dan kekuatan hubungan linear antarvariabel")
        nums=df_model[[c for c in ["Price"]+ANALYSIS_FEATURES if c in df_model.columns]].select_dtypes(include=np.number)
        corr=nums.corr(numeric_only=True); labels=[display_name(c) for c in corr.columns]
        fig=px.imshow(corr,x=labels,y=labels,text_auto=".2f",aspect="auto",color_continuous_scale=[[0,"#eeeeee"],[.5,"#888888"],[1,"#171717"]]); fig.update_coloraxes(showscale=False)
        show_chart(plot_layout(fig,max(500,min(720,34*len(labels)))), "Korelasi Antarvariabel Numerik", "dash_heatmap")
        price_corr=corr["Price"].drop("Price").sort_values(key=lambda x:x.abs(),ascending=False)
        strongest=price_corr.index[0] if len(price_corr) else "-"; strongest_val=price_corr.iloc[0] if len(price_corr) else np.nan
        heat_dir = "positif" if strongest_val >= 0 else "negatif"
        render_insight((f"Hubungan linear numerik terkuat dengan Price adalah {display_name(str(strongest))}, dengan Pearson r = {strongest_val:.2f}.", f"Arah hubungan tersebut cenderung {heat_dir}; Price cenderung bergerak {('naik' if strongest_val >= 0 else 'turun')} ketika predictor meningkat secara linear.", "Temuan ini menjadi gambaran awal sebelum seluruh predictor dianalisis simultan. Korelasi bivariate bukan bukti sebab-akibat dan tidak menggantikan multiple regression."))

        st.markdown('<div style="height:18px"></div>', unsafe_allow_html=True)
        section_head("Mean Price Heatmap", "Rata-rata Harga Mobil berdasarkan Jenis Bahan Bakar dan Jumlah Pintu")
        work = df_model.copy()
        work["Doors"] = pd.to_numeric(work["Doors"], errors="coerce")
        work["Price"] = pd.to_numeric(work["Price"], errors="coerce")
        work = work.dropna(subset=["Fuel Type", "Doors", "Price"])
        mean_pivot = work.pivot_table(index="Fuel Type", columns="Doors", values="Price", aggfunc="mean").sort_index()
        if not mean_pivot.empty:
            mean_pivot.columns = [str(int(c)) for c in mean_pivot.columns]
            mean_fig = px.imshow(mean_pivot, text_auto=".0f", aspect="auto", color_continuous_scale=[[0,"#f2f2f2"],[.5,"#9a9a9a"],[1,"#171717"]])
            mean_fig.update_coloraxes(colorbar_title="Rata-rata Harga (€)")
            mean_fig.update_xaxes(title="Doors")
            mean_fig.update_yaxes(title="Fuel Type")
            show_chart(plot_layout(mean_fig, max(360, 80 + 50 * len(mean_pivot.index))), "Rata-rata Harga Berdasarkan Fuel Type dan Doors", "dash_mean_heatmap_combined")
            render_insight(chart_interpretation_mean_heatmap(df_model))
        else:
            st.info("Data Fuel Type, Doors, dan Price tidak cukup untuk membuat heatmap rata-rata.")
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[4]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Actual vs Predicted Price", "Scatterplot evaluasi model pada data test; garis diagonal menunjukkan prediksi ideal")
        actual_pred = pd.DataFrame({"Harga Aktual": final_y_test.values, "Harga Prediksi": final_y_pred_test.values})
        fig = px.scatter(actual_pred, x="Harga Aktual", y="Harga Prediksi", opacity=.62)
        fig.update_traces(marker=dict(size=6, color="#222222"))
        lo = min(float(actual_pred["Harga Aktual"].min()), float(actual_pred["Harga Prediksi"].min()))
        hi = max(float(actual_pred["Harga Aktual"].max()), float(actual_pred["Harga Prediksi"].max()))
        fig.add_trace(go.Scatter(x=[lo, hi], y=[lo, hi], mode="lines", line=dict(color="#666666", width=2), name="Prediksi ideal"))
        fig.update_xaxes(title="Harga Aktual (€)")
        fig.update_yaxes(title="Harga Prediksi (€)")
        show_chart(plot_layout(fig, 470), "Harga Aktual vs Harga Prediksi", "dash_actual_predicted")
        render_insight((
            f"Sebagian besar titik berada di sekitar garis prediksi ideal; jarak titik dari garis menunjukkan besar error masing-masing observasi.",
            f"Titik di atas garis berarti model memprediksi lebih tinggi daripada harga aktual, sedangkan titik di bawah garis berarti model memprediksi lebih rendah.",
            "Plot ini melihat performa prediksi secara langsung. Untuk membaca apakah error memiliki pola sistematis, gunakan subpage Residuals."
        ))
        st.markdown('</div>', unsafe_allow_html=True)

        st.markdown('<div style="height:18px"></div>', unsafe_allow_html=True)
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Scatterplot per Variabel", "Pilih satu variabel numerik untuk melihat hubungannya dengan Harga Mobil")
        numeric_features = [c for c in ANALYSIS_FEATURES if c in df_model.columns and pd.api.types.is_numeric_dtype(df_model[c])]
        if numeric_features:
            chosen_scatter = st.selectbox(
                "Variabel",
                numeric_features,
                format_func=display_name,
                key="dash_scatter_var",
            )
            plot_df = df_model[[chosen_scatter, "Price"]].copy()
            plot_df[chosen_scatter] = pd.to_numeric(plot_df[chosen_scatter], errors="coerce")
            plot_df["Price"] = pd.to_numeric(plot_df["Price"], errors="coerce")
            plot_df = plot_df.dropna()
            if len(plot_df) > 1 and plot_df[chosen_scatter].nunique() > 1:
                var_fig = px.scatter(plot_df, x=chosen_scatter, y="Price", trendline="ols", opacity=.55)
                var_fig.update_traces(marker=dict(size=6, color="#222222"))
                var_fig.update_xaxes(title=display_name(chosen_scatter))
                var_fig.update_yaxes(title="Harga Mobil (€)")
                show_chart(plot_layout(var_fig, 470), f"{display_name(chosen_scatter)} terhadap Harga Mobil", "dash_scatter_variable")
                render_insight(chart_interpretation_scatter(plot_df, chosen_scatter))
            else:
                st.info(f"Data {display_name(chosen_scatter)} tidak cukup untuk scatterplot.")
        else:
            st.info("Tidak ada variabel numerik yang tersedia untuk scatterplot.")
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[5]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Pola Residual Model", "Memeriksa pola kesalahan prediksi pada data test")
        residual_df=pd.DataFrame({"Harga Prediksi":final_y_pred_test.values,"Residual":final_y_test.values-final_y_pred_test.values})
        fig=px.scatter(residual_df,x="Harga Prediksi",y="Residual",opacity=.62); fig.update_traces(marker=dict(size=5,color="#171717")); fig.add_hline(y=0,line_dash="dash",line_color="#555555")
        show_chart(plot_layout(fig,460), "Residual terhadap Harga Prediksi", "dash_residual")
        resid = residual_df["Residual"]
        corr_resid_pred = float(pd.Series(residual_df["Harga Prediksi"]).corr(pd.Series(resid))) if len(resid)>1 else np.nan
        pattern_note = "tidak menunjukkan hubungan linear yang kuat" if (np.isnan(corr_resid_pred) or abs(corr_resid_pred) < 0.20) else "masih menunjukkan hubungan linear yang perlu diperhatikan"
        render_insight((f"Residual memiliki median {resid.median():,.0f}; MAE test = €{final_metrics_test['mae']:,.0f} menjadi pembanding besar error tipikal.", f"Hubungan residual dengan harga prediksi {pattern_note} (r = {corr_resid_pred:.2f} bila tersedia).", "Residual plot digunakan untuk mencari pola error yang belum tertangkap model. Bentuk kipas, lengkungan, atau kelompok residual yang sistematis perlu diperiksa lebih lanjut."))
        st.markdown('</div>', unsafe_allow_html=True)

    with dash_tabs[6]:
        st.markdown('<div class="card">', unsafe_allow_html=True)
        section_head("Extreme Value Analysis", "Deteksi P1/P99 pada tujuh variabel numerik")
        st.dataframe(extreme_summary.style.set_table_styles([{ "selector":"th", "props":[("background-color","#252627"),("color","#ffffff"),("font-weight","700")] }]).format({"P1 (1%)":"{:,.2f}","P99 (99%)":"{:,.2f}","Lower Extreme (<P1)":"{:,.0f}","Upper Extreme (>P99)":"{:,.0f}","Total Extreme":"{:,.0f}"}),use_container_width=True,hide_index=True)
        render_insight((f"Batas P1/P99 menandai {int(extreme_mask.sum()):,} observasi unik sebagai extreme pada minimal satu variabel.", "Observasi tersebut berada di bagian paling rendah atau paling tinggi dari distribusi variabel yang dianalisis; satu observasi dapat extreme pada lebih dari satu variabel.", "Extreme value tidak otomatis berarti data salah. Karena itu dashboard menyediakan Keep vs Remove untuk melihat dampaknya terhadap error dan kemampuan prediksi model."))
        if not extreme_observations.empty:
            st.dataframe(extreme_observations.style.set_table_styles([{ "selector":"th", "props":[("background-color","#252627"),("color","#ffffff"),("font-weight","700")] }]).format({"Nilai":"{:,.2f}","P1":"{:,.2f}","P99":"{:,.2f}"}),use_container_width=True,hide_index=True,height=380)
        st.markdown('</div>', unsafe_allow_html=True)


# PAGE DISPATCH
# ─────────────────────────────────────────────────────────────────────────────
elif page == "Find & Predict Car Price":
    st.markdown('<div class="eyebrow">Buyer workspace</div><div class="page-title">Find & Predict Car Price</div><div class="page-copy">Satu workspace untuk memilih spesifikasi kendaraan, mengestimasi harga, dan melihat kendaraan yang benar-benar match dengan kriteria tersebut.</div>', unsafe_allow_html=True)
    st.markdown('<div class="rule"></div>', unsafe_allow_html=True)
    render_predict_price()

st.markdown('<div style="text-align:center;color:#9a9a94;font-size:10px;margin-top:45px;letter-spacing:.04em">PRICE ANALYTICS · MULTIPLE LINEAR REGRESSION</div>', unsafe_allow_html=True)
